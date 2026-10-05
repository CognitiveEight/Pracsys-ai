"""
File handlers for processing different file types using GCP Vision OCR.
"""

import base64
import os
import re
import subprocess
import tempfile
from io import BytesIO
from typing import Any, Dict, List, Optional

from loguru import logger

# Known legal-database / publisher watermark patterns that appear on every page
# of exports from SCC Online (TruePrint), Manupatra, Indian Kanoon, etc. These
# supplement the generic cross-page repeated-line detection in
# ``_strip_repeated_pdf_boilerplate`` so the publisher header/footer never
# dominates extracted content (which previously caused image-only judgments to
# skip OCR and be rejected as "repetitive headers").
# Minimum length for a repeated line to be treated as a running header/footer.
# Below this, a recurring line is far more likely to be a body word emitted on
# its own line by the PDF text layer than a header. See _has_header_substance.
_MIN_BOILERPLATE_LINE_CHARS = 12

_PDF_BOILERPLATE_PATTERNS = [
    re.compile(r"scc\s*online", re.I),
    re.compile(r"trueprint", re.I),
    re.compile(r"printed\s+for\s*:", re.I),
    re.compile(r"printed\s+by\s*:", re.I),
    re.compile(r"scconline\.com", re.I),
    re.compile(r"eastern\s+book\s+company", re.I),
    re.compile(r"ebc\s+publishing", re.I),
    re.compile(r"manupatra", re.I),
    re.compile(r"indiankanoon", re.I),
    re.compile(r"downloaded\s+on", re.I),
    re.compile(r"^\s*page\s+\d+\s*$", re.I),
    re.compile(r"^\s*-{20,}\s*$"),
    re.compile(r"^\s*_{20,}\s*$"),
    # Copyright citation line, e.g. "... v. D.B. Modak, (2008) 1 SCC 1 ..."
    re.compile(r"\(20\d\d\)\s+\d+\s+scc\s+\d+", re.I),
]


def _strip_repeated_pdf_boilerplate(
    page_texts: List[str],
    *,
    page_numbers: Optional[List[int]] = None,
    mark_pages: bool = False,
) -> str:
    """Remove per-page publisher header/footer boilerplate from extracted PDF text.

    Two signals, combined:
      1. Cross-page repetition (publisher-agnostic): a normalized non-empty line
         that recurs on ``>= max(3, 0.4 * num_pages)`` pages is treated as a
         running header/footer and dropped.
      2. Known legal-database watermark patterns (``_PDF_BOILERPLATE_PATTERNS``)
         are always dropped, so even a single-page export is cleaned.

    Returns the cleaned text (pages rejoined). Deterministic and idempotent.
    Only whole recurring lines are removed — never paragraphs — so legitimate
    body text is preserved.
    """
    pages = [p or "" for p in (page_texts or [])]
    num_pages = len(pages)
    if num_pages == 0:
        return ""

    def _norm(line: str) -> str:
        return re.sub(r"\s+", " ", line.strip()).lower()

    # Count how many pages each normalized line appears on.
    line_page_count: Dict[str, int] = {}
    for page in pages:
        seen_on_page = set()
        for raw_line in page.splitlines():
            norm = _norm(raw_line)
            if not norm or norm in seen_on_page:
                continue
            seen_on_page.add(norm)
            line_page_count[norm] = line_page_count.get(norm, 0) + 1

    repeat_threshold = max(3, int(0.4 * num_pages))

    def _has_header_substance(norm: str) -> bool:
        """Could this line plausibly BE a running header/footer?

        The cross-page-repetition signal assumes a line is a whole header or
        footer. That assumption collapses when the PDF's text layer comes out
        one word per line — which PyPDF2 does for many documents. Ordinary
        function words ("the", "and", "of", "to", "is") then recur on every
        page and get deleted as boilerplate, silently gutting the body text:
        "our mission is to improve the effectiveness and sophistication of
        partnership" became "our mission improve effectiveness sophistication
        partnership". The model is then left to reconstruct the missing
        connective tissue, which is where invented facts come from.

        A genuine running header carries substance — a publisher name, a case
        citation, a page marker. Require that before the repetition signal may
        fire. Known watermark patterns are matched separately and are not
        subject to this guard.
        """
        # A SINGLE word is never a running header, however long: with one
        # word per line, "verification" (12 chars) recurring on 3 of 4 pages
        # was deleted, turning "BCCI Age Verification Committee" into
        # "BCCI Age Committee" in every panel. Publisher one-word marks
        # ("Manupatra") are caught by _PDF_BOILERPLATE_PATTERNS instead.
        words = len(norm.split())
        return words >= 3 or (words >= 2 and len(norm) >= _MIN_BOILERPLATE_LINE_CHARS)

    def _is_boilerplate(raw_line: str) -> bool:
        norm = _norm(raw_line)
        if not norm:
            return False
        if (
            line_page_count.get(norm, 0) >= repeat_threshold
            and _has_header_substance(norm)
        ):
            return True
        return any(pat.search(raw_line) for pat in _PDF_BOILERPLATE_PATTERNS)

    cleaned_pages: List[str] = []
    marked_pages: List[str] = []
    for pos, page in enumerate(pages):
        kept = [ln for ln in page.splitlines() if not _is_boilerplate(ln)]
        page_body = "\n".join(kept).strip()
        if page_body:
            cleaned_pages.append(page_body)
            number = (
                page_numbers[pos]
                if page_numbers and pos < len(page_numbers)
                else pos + 1
            )
            marked_pages.append(f"--- Page {number} ---\n{page_body}")
    cleaned = "\n".join(cleaned_pages).strip()

    # A header/footer strip should shave a few percent. Much more than that
    # means the repetition heuristic is eating body text, which corrupts every
    # downstream prompt — surface it instead of failing silently.
    original_len = sum(len(p) for p in pages)
    if original_len and (original_len - len(cleaned)) / original_len > 0.25:
        logger.warning(
            "PDF boilerplate strip removed {:.0f}% of the text layer ({} -> {} "
            "chars over {} pages) — the repetition heuristic may be eating body "
            "text; check the PDF's text layer",
            100 * (original_len - len(cleaned)) / original_len,
            original_len,
            len(cleaned),
            num_pages,
        )
    if mark_pages:
        # Physical page numbers, so a citation can say "page 7" and be right.
        # Blank pages are skipped but keep their numbering.
        return "\n\n".join(marked_pages).strip()
    return cleaned

try:
    from google.auth.exceptions import DefaultCredentialsError
    from google.cloud import vision

    GCP_OCR_AVAILABLE = True
except ImportError:
    GCP_OCR_AVAILABLE = False
    logger.warning("Google Cloud Vision API not available")

try:
    import PyPDF2

    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    logger.warning("PyPDF2 not available")

try:
    from pdf2image import convert_from_bytes

    PDF2IMAGE_AVAILABLE = True
except ImportError:
    PDF2IMAGE_AVAILABLE = False
    logger.warning("pdf2image not available")

try:
    from PIL import Image

    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False
    logger.warning("PIL/Pillow not available")

# Max pixels per image for OCR (avoids DecompressionBomb and GCP Vision "Bad image data")
# PIL default is ~89M; GCP Vision has size limits; keep safe for memory and API
OCR_MAX_IMAGE_PIXELS = 20_000_000

try:
    from docx import Document

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not available")

try:
    import docx2txt

    DOC_AVAILABLE = True
except ImportError:
    DOC_AVAILABLE = False
    logger.warning("docx2txt not available")

try:
    import openpyxl

    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False
    logger.warning("openpyxl not available")

try:
    from pptx import Presentation

    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False
    logger.warning("python-pptx not available")


AUDIO_EXTENSIONS = (
    ".mp3", ".wav", ".m4a", ".ogg", ".flac", ".opus", ".webm", ".aac",
)

XLSX_EXTENSIONS = (".xlsx", ".xlsm")
PPTX_EXTENSIONS = (".pptx",)
# Legacy binary Office formats. Without a dedicated parser these used to fall
# into the OLE2 sniff and come out of the DOC extractor as raw-byte noise; they
# are now rejected explicitly so the readability score reports a failure.
LEGACY_OFFICE_EXTENSIONS = (".xls", ".ppt")
_IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".webp")
#: Every extension routed by name in ``process_file_by_extension``; anything
#: else is identified from its content.
KNOWN_EXTENSIONS = frozenset(
    AUDIO_EXTENSIONS + XLSX_EXTENSIONS + PPTX_EXTENSIONS + LEGACY_OFFICE_EXTENSIONS
    + _IMAGE_EXTENSIONS + (".pdf", ".docx", ".doc", ".txt", ".md", ".markdown")
)

# A slide with less text than this of its own has its pictures OCR'd (it is
# most likely a scanned page pasted onto a slide). The cap bounds Vision spend
# on a photo-heavy deck; each OCR'd picture is billed as one ocrUnit.
PPTX_OCR_TEXT_THRESHOLD = 20
PPTX_MAX_OCR_IMAGES = 50


def filename_for_extraction(name: Optional[str], extension: Optional[str]) -> str:
    """Return ``name`` with ``extension`` appended when it is missing.

    The backend's display ``filename`` ("annexure-3-age-verification-test-results")
    carries no suffix; the ``extension`` field is the only type signal. Use the
    result for ROUTING only (``process_file_by_extension``, ``is_audio_file``)
    and keep the original name for storage and display.
    """
    name = (name or "").strip()
    ext = (extension or "").strip().lstrip(".").lower()
    if ext and not name.lower().endswith("." + ext):
        return f"{name}.{ext}"
    return name


def _zip_office_kind(data: bytes) -> str:
    """Identify a ZIP-based Office file by its parts: ``xlsx``/``pptx``/``docx``.

    Falls back to ``docx``, which is what every ``PK`` payload was treated as
    before spreadsheets and decks were supported.
    """
    import zipfile

    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            names = set(zf.namelist())
    except (zipfile.BadZipFile, ValueError):
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    return "docx"


class FileProcessor:
    """Utility class to process different file types and extract text content."""

    @staticmethod
    def is_url(content: str) -> bool:
        """Check if content is a URL."""
        return content.startswith(("http://", "https://"))

    @staticmethod
    def is_base64_image(content: str) -> bool:
        """Check if content is a base64 encoded image."""
        return content.startswith("data:image/") and ";base64," in content

    @staticmethod
    def is_audio_file(filename: str) -> bool:
        """Return True when the filename has an audio extension we transcribe."""
        if not filename:
            return False
        return filename.lower().endswith(AUDIO_EXTENSIONS)

    @staticmethod
    def is_spreadsheet(filename: str, data: bytes = b"") -> bool:
        """True for an XLSX/XLSM workbook: by extension, or by content when the
        name carries no recognised extension (the same sniff
        ``process_file_by_extension`` falls back to)."""
        name = (filename or "").lower()
        if name.endswith(XLSX_EXTENSIONS):
            return True
        if os.path.splitext(name)[1] in KNOWN_EXTENSIONS:
            return False
        return bool(data) and data.startswith(b"PK") and _zip_office_kind(data) == "xlsx"

    @staticmethod
    def extract_text_from_image(image_data: bytes) -> str:
        """Extract text from image using GCP Vision API OCR."""
        if not GCP_OCR_AVAILABLE:
            logger.warning(
                "GCP Vision API not available. Cannot extract text from image."
            )
            return "[Image content - GCP Vision API not available]"

        try:
            import os

            from google.auth.api_key import Credentials as ApiKeyCredentials

            from app.ai.usage import record_ocr_units
            from app.core.config import settings

            # Vision OCR uses API key only (no service-account credentials).
            api_key_value = (
                (getattr(settings, "google_cloud_vision_api_key", None) or "")
                or os.getenv("GOOGLE_CLOUD_VISION_API_KEY")
                or ""
            ).strip()
            if not api_key_value:
                logger.error(
                    "Google Vision API key not configured. Set GOOGLE_CLOUD_VISION_API_KEY."
                )
                return "[Error: Vision API key not configured (set GOOGLE_CLOUD_VISION_API_KEY)]"

            api_credentials = ApiKeyCredentials(api_key_value)
            client = vision.ImageAnnotatorClient(credentials=api_credentials)

            image = vision.Image(content=image_data)
            response = client.text_detection(image=image)
            # Cost: one billable Vision unit per call — recorded here, the single
            # Vision call site, so every OCR path (per-PDF-page loop in
            # _extract_text_from_pdf_with_ocr, standalone images, content-sniffed
            # images) is counted exactly once. Recorded BEFORE the error check
            # below because a response that carries an error was still billed.
            record_ocr_units(1)
            if response.error.message:
                raise Exception(f"GCP Vision API error: {response.error.message}")

            texts = response.text_annotations
            if texts:
                extracted_text = texts[0].description
                logger.success(
                    "Extracted {} characters from image using API key", len(extracted_text)
                )
                return extracted_text.strip()

            logger.warning("No text detected in image")
            return "[No text detected in image]"

        except Exception as e:
            logger.error("Error extracting text from image: {}", str(e))
            return f"[Error extracting text from image: {str(e)}]"

    # Service-account OCR auth helpers removed (API key only).

    @staticmethod
    def extract_text_from_pdf(pdf_data: bytes, mark_pages: bool = False) -> str:
        """Extract text from PDF file using text extraction and OCR for image-based PDFs."""
        if not PDF_AVAILABLE:
            logger.warning(
                "PDF processing library not available. Cannot extract text from PDF."
            )
            return "[PDF content - PDF processing not available]"

        try:
            # First try traditional text extraction, keeping per-page text so we
            # can detect (and strip) repeated publisher header/footer boilerplate.
            pdf_file = BytesIO(pdf_data)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            page_texts = []

            for page in pdf_reader.pages:
                page_text = page.extract_text()
                page_texts.append(page_text or "")

            num_pages = max(1, len(page_texts))
            # Strip per-page publisher watermark so a boilerplate-only text layer
            # (e.g. an SCC Online / TruePrint image scan whose only text is the
            # repeated header) does not masquerade as real content and skip OCR.
            scrubbed = _strip_repeated_pdf_boilerplate(page_texts)

            # Keep the text layer ONLY when it carries meaningful body: a decent
            # absolute amount AND a reasonable average of real chars per page.
            # Otherwise the text layer is boilerplate-dominated / image-only and
            # we must OCR the page images to read the actual document.
            has_real_text = (
                len(scrubbed) >= 200 and (len(scrubbed) / num_pages) >= 40
            )
            if has_real_text:
                logger.success(
                    "Extracted {} characters from PDF using text extraction "
                    "({} chars after boilerplate strip, {} pages)",
                    sum(len(t) for t in page_texts),
                    len(scrubbed),
                    num_pages,
                )
                # The OCR decision above is made on the UNMARKED text: markers
                # add ~15 chars per page and would tip a boilerplate-only
                # document over the "real text" threshold.
                if mark_pages:
                    return _strip_repeated_pdf_boilerplate(page_texts, mark_pages=True)
                return scrubbed

            # Boilerplate-dominated / image-based → OCR the page images.
            if PDF2IMAGE_AVAILABLE and GCP_OCR_AVAILABLE:
                logger.info(
                    "Text layer is boilerplate-dominated ({} real chars over {} "
                    "pages); attempting OCR on PDF pages",
                    len(scrubbed),
                    num_pages,
                )
                return FileProcessor._extract_text_from_pdf_with_ocr(
                    pdf_data, mark_pages=mark_pages
                )
            else:
                logger.warning(
                    "PDF appears to be image-based but pdf2image or GCP Vision API not available"
                )
                return (
                    scrubbed
                    if scrubbed
                    else "[PDF appears to contain images but OCR not available]"
                )

        except Exception as e:
            logger.error("Error extracting text from PDF: {}", str(e))
            return f"[Error extracting text from PDF: {str(e)}]"

    @staticmethod
    def _extract_text_from_pdf_with_ocr(pdf_data: bytes, mark_pages: bool = False) -> str:
        """Extract text from image-based PDF using OCR."""
        try:
            resample = None
            old_max_pixels = None
            if PIL_AVAILABLE:
                resample = Image.LANCZOS if hasattr(Image, "LANCZOS") else Image.BICUBIC
                if hasattr(Image, "Resampling") and hasattr(
                    Image.Resampling, "LANCZOS"
                ):
                    resample = Image.Resampling.LANCZOS
                old_max_pixels = getattr(Image, "MAX_IMAGE_PIXELS", None)
                if old_max_pixels is not None:
                    Image.MAX_IMAGE_PIXELS = 200_000_000
            try:
                # Convert at 150 DPI to reduce size (default 200 can exceed limits)
                images = convert_from_bytes(pdf_data, dpi=150)
            finally:
                if PIL_AVAILABLE and old_max_pixels is not None:
                    Image.MAX_IMAGE_PIXELS = old_max_pixels

            extracted_texts = []
            extracted_page_numbers: List[int] = []

            for i, image in enumerate(images):
                logger.info("Processing PDF page {}/{} with OCR", i + 1, len(images))

                # Resize if image exceeds safe pixel count (avoids GCP Vision "Bad image data")
                if (
                    PIL_AVAILABLE
                    and resample is not None
                    and getattr(image, "width", None)
                    and getattr(image, "height", None)
                ):
                    w, h = image.width, image.height
                    pixels = w * h
                    if pixels > OCR_MAX_IMAGE_PIXELS:
                        ratio = (OCR_MAX_IMAGE_PIXELS / pixels) ** 0.5
                        new_w = max(1, int(w * ratio))
                        new_h = max(1, int(h * ratio))
                        image = image.resize((new_w, new_h), resample)
                        logger.warning(
                            "PDF page %s image too large (%s pixels); resized to %sx%s for OCR",
                            i + 1,
                            pixels,
                            new_w,
                            new_h,
                        )

                # Convert PIL image to bytes
                img_byte_arr = BytesIO()
                image.save(img_byte_arr, format="PNG")
                img_bytes = img_byte_arr.getvalue()

                # Extract text using GCP Vision API
                page_text = FileProcessor.extract_text_from_image(img_bytes)

                if (
                    page_text
                    and not page_text.startswith("[Error")
                    and not page_text.startswith("[No text")
                ):
                    extracted_texts.append(page_text)
                    extracted_page_numbers.append(i + 1)

            if extracted_texts:
                # OCR also reads the publisher header/footer imprinted on each
                # page image, so strip the repeated boilerplate here too.
                full_text = _strip_repeated_pdf_boilerplate(
                    extracted_texts,
                    page_numbers=extracted_page_numbers,
                    mark_pages=mark_pages,
                )
                if not full_text:
                    full_text = "\n\n".join(extracted_texts)
                logger.success(
                    "Extracted {} characters from PDF using OCR", len(full_text)
                )
                return full_text
            else:
                return "[No text could be extracted from PDF images]"

        except Exception as e:
            logger.error("Error extracting text from PDF with OCR: {}", str(e))
            return f"[Error extracting text from PDF with OCR: {str(e)}]"

    @staticmethod
    def extract_text_from_doc(doc_data: bytes) -> str:
        """Extract text from legacy DOC file."""
        if not DOC_AVAILABLE:
            logger.warning(
                "DOC processing library not available. Cannot extract text from DOC."
            )
            return "[DOC content - DOC processing not available]"

        try:
            # Save bytes to temporary file for docx2txt processing
            with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as temp_file:
                temp_file.write(doc_data)
                temp_file_path = temp_file.name

            try:
                # First try docx2txt
                try:
                    text = docx2txt.process(temp_file_path)
                    if text and text.strip():
                        logger.success("Extracted {} characters from DOC", len(text))
                        return text.strip()
                except Exception as e:
                    logger.warning(
                        "docx2txt failed: {}, trying alternative method", str(e)
                    )

                # Fallback: try using antiword command if available
                try:
                    result = subprocess.run(
                        ["antiword", temp_file_path],
                        capture_output=True,
                        text=True,
                        timeout=30,
                    )
                    if result.returncode == 0 and result.stdout.strip():
                        logger.success(
                            "Extracted {} characters from DOC using antiword", len(result.stdout)
                        )
                        return result.stdout.strip()
                except (
                    subprocess.TimeoutExpired,
                    subprocess.CalledProcessError,
                    FileNotFoundError,
                ):
                    logger.warning("antiword command not available or failed")

                # Final fallback: try to extract as raw text (very basic)
                try:
                    with open(temp_file_path, "rb") as f:
                        content = f.read()

                    # Look for text patterns in the binary data
                    text_parts = []
                    i = 0
                    while i < len(content) - 1:
                        if content[i] >= 32 and content[i] <= 126:  # Printable ASCII
                            start = i
                            while (
                                i < len(content)
                                and content[i] >= 32
                                and content[i] <= 126
                            ):
                                i += 1
                            if (
                                i - start > 3
                            ):  # Only consider sequences of 4+ characters
                                text_parts.append(
                                    content[start:i].decode("ascii", errors="ignore")
                                )
                        i += 1

                    if text_parts:
                        extracted_text = " ".join(text_parts)
                        # Clean up the extracted text
                        cleaned_text = " ".join(extracted_text.split())
                        if (
                            len(cleaned_text) > 10
                        ):  # Only return if we got substantial text
                            logger.success(
                                "Extracted {} characters from DOC using raw extraction", len(cleaned_text)
                            )
                            return cleaned_text

                except Exception as e:
                    logger.warning("Raw text extraction failed: {}", str(e))

                return "[Could not extract text from DOC file - format may be corrupted or unsupported]"

            finally:
                # Clean up temporary file
                try:
                    os.unlink(temp_file_path)
                except:
                    pass

        except Exception as e:
            logger.error("Error extracting text from DOC: {}", str(e))
            return f"[Error extracting text from DOC: {str(e)}]"

    @staticmethod
    def extract_text_from_docx(docx_data: bytes) -> str:
        """Extract text from DOCX file.

        Word auto-numbered list items (``<w:numPr>``) carry their numbers in
        ``numbering.xml``, NOT in ``paragraph.text`` — so a naive extraction
        drops the "1. 2. 3." an advocate sees on a Vakalatnama's authorisation
        clauses, a petition's grounds, or a prayer list. We resolve each
        numbered paragraph to its visible marker (decimal / a) / i) / bullet)
        and prefix it, so the downstream drafter sees the list structure and
        can reproduce it. ``numId == 0`` means "no list" (numbering removed)
        and is left unprefixed.
        """
        if not DOCX_AVAILABLE:
            logger.warning(
                "DOCX processing library not available. Cannot extract text from DOCX."
            )
            return "[DOCX content - DOCX processing not available]"

        try:
            docx_file = BytesIO(docx_data)
            doc = Document(docx_file)

            numfmt_map = FileProcessor._docx_numbering_formats(doc)
            counters: dict = {}
            lines: list = []
            for paragraph in doc.paragraphs:
                marker = FileProcessor._docx_list_marker(
                    paragraph, numfmt_map, counters
                )
                txt = paragraph.text
                lines.append(f"{marker}{txt}" if marker else txt)

            text = "\n".join(lines)
            logger.success("Extracted {} characters from DOCX", len(text))
            return text.strip()
        except Exception as e:
            logger.error("Error extracting text from DOCX: {}", str(e))
            return f"[Error extracting text from DOCX: {str(e)}]"

    @staticmethod
    def _docx_numbering_formats(doc) -> dict:
        """Map ``numId -> {ilvl: (numFmt, start)}`` from the document's
        numbering part. Returns ``{}`` when the document has no numbering
        part (no auto-numbered lists)."""
        from docx.oxml.ns import qn

        try:
            numbering = doc.part.numbering_part.element
        except Exception:
            return {}

        # abstractNumId -> {ilvl: (numFmt, start)}
        abstract: dict = {}
        for anum in numbering.findall(qn("w:abstractNum")):
            aid = anum.get(qn("w:abstractNumId"))
            levels: dict = {}
            for lvl in anum.findall(qn("w:lvl")):
                try:
                    ilvl = int(lvl.get(qn("w:ilvl")) or 0)
                except (TypeError, ValueError):
                    ilvl = 0
                fmt_el = lvl.find(qn("w:numFmt"))
                fmt = fmt_el.get(qn("w:val")) if fmt_el is not None else "decimal"
                start_el = lvl.find(qn("w:start"))
                try:
                    start = int(start_el.get(qn("w:val"))) if start_el is not None else 1
                except (TypeError, ValueError):
                    start = 1
                levels[ilvl] = (fmt, start)
            if aid is not None:
                abstract[aid] = levels

        num_map: dict = {}
        for num in numbering.findall(qn("w:num")):
            nid = num.get(qn("w:numId"))
            aref = num.find(qn("w:abstractNumId"))
            aid = aref.get(qn("w:val")) if aref is not None else None
            if nid is not None and aid in abstract:
                num_map[nid] = abstract[aid]
        return num_map

    @staticmethod
    def _docx_list_marker(paragraph, numfmt_map: dict, counters: dict) -> str:
        """Return the visible list marker (e.g. ``"1. "``, ``"a) "``, ``"• "``)
        for a numbered paragraph, advancing the per-(numId, ilvl) counter, or
        ``""`` when the paragraph is not part of a real auto-numbered list."""
        from docx.oxml.ns import qn

        pPr = paragraph._p.pPr
        if pPr is None:
            return ""
        numPr = pPr.find(qn("w:numPr"))
        if numPr is None:
            return ""
        num_id_el = numPr.find(qn("w:numId"))
        ilvl_el = numPr.find(qn("w:ilvl"))
        num_id = num_id_el.get(qn("w:val")) if num_id_el is not None else None
        # numId 0 (or absent) means numbering was removed for this paragraph.
        if not num_id or num_id == "0":
            return ""
        try:
            ilvl = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
        except (TypeError, ValueError):
            ilvl = 0

        levels = numfmt_map.get(num_id)
        if not levels or ilvl not in levels:
            # Unknown list definition — fall back to a generic decimal marker
            # so the structure is still visible.
            fmt, start = "decimal", 1
        else:
            fmt, start = levels[ilvl]

        if fmt == "bullet":
            return ("  " * ilvl) + "• "

        key = (num_id, ilvl)
        counters[key] = counters.get(key, start - 1) + 1
        # Reset any deeper levels under the same list when a higher level ticks.
        for k in [k for k in counters if k[0] == num_id and k[1] > ilvl]:
            del counters[k]
        n = counters[key]

        return ("  " * ilvl) + FileProcessor._format_list_number(fmt, n) + " "

    @staticmethod
    def _format_list_number(fmt: str, n: int) -> str:
        """Format an ordinal ``n`` per a Word numFmt into its marker text."""
        def to_roman(num: int) -> str:
            vals = [
                (1000, "m"), (900, "cm"), (500, "d"), (400, "cd"),
                (100, "c"), (90, "xc"), (50, "l"), (40, "xl"),
                (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"),
            ]
            out = []
            for v, sym in vals:
                while num >= v:
                    out.append(sym)
                    num -= v
            return "".join(out)

        def to_letter(num: int) -> str:
            # 1->a, 2->b, ... 26->z, 27->aa (Excel-style)
            s = ""
            while num > 0:
                num, rem = divmod(num - 1, 26)
                s = chr(ord("a") + rem) + s
            return s

        if fmt == "lowerLetter":
            return f"{to_letter(n)})"
        if fmt == "upperLetter":
            return f"{to_letter(n).upper()})"
        if fmt == "lowerRoman":
            return f"{to_roman(n)}."
        if fmt == "upperRoman":
            return f"{to_roman(n).upper()}."
        # decimal and anything unrecognised
        return f"{n}."

    @staticmethod
    def extract_text_from_xlsx(xlsx_data: bytes) -> str:
        """Extract text from an XLSX/XLSM workbook.

        Parsing and rendering live in ``app/processing/tables.py``: a workbook
        summary computed from every row, then the rows in blocks that each
        restate the sheet, the row range and the header. Evidence ingestion
        does not come through here; ``DocumentProcessor`` uses the parsed
        tables directly so it can redact by column and store the rows. This
        text-only path serves chat attachments and content-sniffed files.
        """
        if not XLSX_AVAILABLE:
            logger.warning(
                "XLSX processing library not available. Cannot extract text from XLSX."
            )
            return "[XLSX content - XLSX processing not available]"
        try:
            from app.processing.tables import extract_text

            return extract_text(xlsx_data).strip()
        except Exception as e:
            logger.error("Error extracting text from XLSX: {}", str(e))
            return f"[Error extracting text from XLSX: {str(e)}]"

    @staticmethod
    def _pptx_shape_lines(shape, lines: List[str], pictures: List[Any]) -> None:
        """Append one shape's text to ``lines``, recursing into groups.

        Picture shapes are collected into ``pictures`` for optional OCR.
        """
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in FileProcessor._pptx_reading_order(shape.shapes):
                FileProcessor._pptx_shape_lines(child, lines, pictures)
            return
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pictures.append(shape)
            return
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [
                    re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells
                ]
                if any(cells):
                    lines.append(" | ".join(cells))
            return
        if getattr(shape, "has_chart", False) and shape.has_chart:
            FileProcessor._pptx_chart_lines(shape.chart, lines)
            return
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            # python-pptx renders a soft line break (<a:br>) as a vertical tab.
            text = shape.text_frame.text.replace("\v", "\n").strip()
            if text:
                lines.append(text)

    @staticmethod
    def _pptx_chart_lines(chart, lines: List[str]) -> None:
        """Append a chart's title and its data points as text.

        A chart in a presentation is often the only place a figure appears,
        so its numbers are extracted rather than dropped with the graphic.
        """
        try:
            title = ""
            if chart.has_title and chart.chart_title.has_text_frame:
                title = chart.chart_title.text_frame.text.strip()
            lines.append(f"Chart: {title}" if title else "Chart:")
            for plot in chart.plots:
                categories = [str(c) for c in plot.categories]
                for series in plot.series:
                    from app.processing.tables import make_cell

                    values = [make_cell(v).display for v in series.values]
                    if categories and len(categories) == len(values):
                        points = ", ".join(
                            f"{c} = {v}" for c, v in zip(categories, values)
                        )
                    else:
                        points = ", ".join(values)
                    lines.append(f"{series.name or 'Series'}: {points}")
        except Exception as e:
            # Chart XML varies a lot by producer; a chart that will not parse
            # must not lose the rest of the slide.
            logger.warning("Could not read chart data from PPTX: {}", str(e))

    @staticmethod
    def _pptx_reading_order(shapes) -> list:
        """Order shapes top-to-bottom, then left-to-right.

        ``slide.shapes`` is in z-order (the order they were drawn), which often
        puts a caption before the body it belongs under.
        """
        return sorted(
            shapes, key=lambda s: ((s.top or 0), (s.left or 0))
        )

    @staticmethod
    def extract_text_from_pptx(pptx_data: bytes, mark_pages: bool = False) -> str:
        """Extract text from a PPTX deck, slide by slide.

        Each slide gives its title, then its text boxes, tables and chart data
        in reading order, then its speaker notes. Hidden slides are included
        and labelled. With ``mark_pages`` each slide is preceded by
        ``--- Page N ---``, where N is the slide number. That is the page the
        slide lands on when the deck is exported to PDF, so chunk page ranges
        and citations point at the slide the lawyer sees.

        A slide with almost no text of its own (a scanned page pasted in as a
        picture, which is common for annexures) has its pictures OCR'd through
        ``extract_text_from_image``, so OCR is metered there exactly as it is
        for PDFs.
        """
        if not PPTX_AVAILABLE:
            logger.warning(
                "PPTX processing library not available. Cannot extract text from PPTX."
            )
            return "[PPTX content - PPTX processing not available]"

        try:
            deck = Presentation(BytesIO(pptx_data))
            output: List[str] = []
            ocr_images = 0
            for number, slide in enumerate(deck.slides, start=1):
                title_shape = slide.shapes.title
                title = (
                    title_shape.text_frame.text.replace("\v", " ").strip()
                    if title_shape is not None and title_shape.has_text_frame
                    else ""
                )
                lines: List[str] = []
                pictures: List[Any] = []
                for shape in FileProcessor._pptx_reading_order(slide.shapes):
                    if title_shape is not None and shape.shape_id == title_shape.shape_id:
                        continue
                    FileProcessor._pptx_shape_lines(shape, lines, pictures)

                own_text = len(title) + sum(len(line) for line in lines)
                if pictures and own_text < PPTX_OCR_TEXT_THRESHOLD:
                    for picture in pictures:
                        if ocr_images >= PPTX_MAX_OCR_IMAGES:
                            break
                        try:
                            blob = picture.image.blob
                        except Exception:
                            continue  # linked, not embedded: nothing to read
                        ocr_images += 1
                        ocr_text = FileProcessor.extract_text_from_image(blob)
                        if ocr_text and not ocr_text.startswith("["):
                            lines.append(ocr_text.strip())

                notes = ""
                if slide.has_notes_slide:
                    notes_frame = slide.notes_slide.notes_text_frame
                    if notes_frame is not None:
                        notes = notes_frame.text.replace("\v", "\n").strip()

                heading = f"Slide {number}: {title}" if title else f"Slide {number}"
                if slide._element.get("show") == "0":
                    heading += " (hidden)"
                parts = [heading, *lines]
                if notes:
                    parts.append(f"Speaker notes: {notes}")
                # Same marker layout as the PDF path: marker, newline, body.
                if mark_pages:
                    parts.insert(0, f"--- Page {number} ---")
                output.append("\n".join(parts))

            text = "\n\n".join(output)
            logger.success(
                "Extracted {} characters from PPTX ({} slides, {} images OCR'd)",
                len(text),
                len(deck.slides),
                ocr_images,
            )
            return text.strip()
        except Exception as e:
            logger.error("Error extracting text from PPTX: {}", str(e))
            return f"[Error extracting text from PPTX: {str(e)}]"

    @staticmethod
    def extract_text_from_txt(txt_data: bytes) -> str:
        """Extract text from TXT file with encoding detection."""
        try:
            # Try UTF-8 first
            try:
                text = txt_data.decode("utf-8")
                logger.success("Extracted {} characters from TXT (UTF-8)", len(text))
                return text.strip()
            except UnicodeDecodeError:
                pass

            # Try other common encodings
            for encoding in ["latin-1", "cp1252", "iso-8859-1"]:
                try:
                    text = txt_data.decode(encoding)
                    logger.success(
                        "Extracted {} characters from TXT ({})", len(text), encoding
                    )
                    return text.strip()
                except UnicodeDecodeError:
                    continue

            # Last resort: decode with errors='ignore'
            text = txt_data.decode("utf-8", errors="ignore")
            logger.warning(
                "Extracted {} characters from TXT with encoding errors ignored", len(text)
            )
            return text.strip()

        except Exception as e:
            logger.error("Error extracting text from TXT: {}", str(e))
            return f"[Error extracting text from TXT: {str(e)}]"

    @staticmethod
    def extract_text_from_md(md_data: bytes) -> str:
        """Extract text from Markdown file."""
        try:
            text = md_data.decode("utf-8", errors="ignore")
            logger.success("Extracted {} characters from Markdown", len(text))
            return text.strip()
        except Exception as e:
            logger.error("Error extracting text from Markdown: {}", str(e))
            return f"[Error extracting text from Markdown: {str(e)}]"

    @staticmethod
    def process_file_by_extension(
        file_data: bytes, filename: str, mark_pages: bool = False
    ) -> str:
        """
        Process file based on its extension.

        Args:
            file_data: File content as bytes
            filename: Name of the file with extension

        Returns:
            Extracted text content
        """
        filename_lower = filename.lower()

        # Audio files: handled separately by document_processor (Chirp 3 transcription).
        # Returning a sentinel here protects callers that route everything through this
        # method but should not try to OCR audio bytes.
        if filename_lower.endswith(AUDIO_EXTENSIONS):
            logger.debug(
                "Audio file routed through file_handlers; document_processor handles transcription"
            )
            return "[Audio file - transcription handled by document_processor]"

        # Spreadsheets and decks. Checked before images: the image test below
        # is a substring match, and a suffix is the stronger signal.
        if filename_lower.endswith(XLSX_EXTENSIONS):
            logger.info("Processing XLSX file: {}", filename)
            return FileProcessor.extract_text_from_xlsx(file_data)

        if filename_lower.endswith(PPTX_EXTENSIONS):
            logger.info("Processing PPTX file: {}", filename)
            return FileProcessor.extract_text_from_pptx(
                file_data, mark_pages=mark_pages
            )

        if filename_lower.endswith(LEGACY_OFFICE_EXTENSIONS):
            logger.warning("Unsupported legacy Office format: {}", filename)
            return (
                "[Unsupported file format - legacy .xls/.ppt files cannot be read; "
                "re-save as .xlsx/.pptx]"
            )

        # Image files
        if any(
            ext in filename_lower
            for ext in [".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".webp"]
        ):
            logger.info("Processing image file: {}", filename)
            return FileProcessor.extract_text_from_image(file_data)

        # PDF files
        elif filename_lower.endswith(".pdf"):
            logger.info("Processing PDF file: {}", filename)
            return FileProcessor.extract_text_from_pdf(file_data, mark_pages=mark_pages)

        # DOCX files
        elif filename_lower.endswith(".docx"):
            logger.info("Processing DOCX file: {}", filename)
            return FileProcessor.extract_text_from_docx(file_data)

        # DOC files
        elif filename_lower.endswith(".doc"):
            logger.info("Processing DOC file: {}", filename)
            return FileProcessor.extract_text_from_doc(file_data)

        # Text files
        elif filename_lower.endswith(".txt"):
            logger.info("Processing TXT file: {}", filename)
            return FileProcessor.extract_text_from_txt(file_data)

        # Markdown files
        elif filename_lower.endswith((".md", ".markdown")):
            logger.info("Processing Markdown file: {}", filename)
            return FileProcessor.extract_text_from_md(file_data)

        # Try to detect by content
        else:
            logger.info("Detecting file type by content for: {}", filename)
            if file_data.startswith(b"%PDF"):
                return FileProcessor.extract_text_from_pdf(
                    file_data, mark_pages=mark_pages
                )
            elif file_data.startswith(b"PK"):  # ZIP-based Office: DOCX/XLSX/PPTX
                kind = _zip_office_kind(file_data)
                if kind == "xlsx":
                    return FileProcessor.extract_text_from_xlsx(file_data)
                if kind == "pptx":
                    return FileProcessor.extract_text_from_pptx(
                        file_data, mark_pages=mark_pages
                    )
                return FileProcessor.extract_text_from_docx(file_data)
            elif file_data.startswith(b"\xd0\xcf\x11\xe0"):  # OLE2 format (DOC files)
                return FileProcessor.extract_text_from_doc(file_data)
            elif file_data.startswith(
                (b"\xff\xd8\xff", b"\x89PNG", b"GIF8")
            ):  # Image files
                return FileProcessor.extract_text_from_image(file_data)
            else:
                # Assume it's a text file
                logger.warning("Unknown file type, treating as text: {}", filename)
                return FileProcessor.extract_text_from_txt(file_data)


# Global file processor instance
file_processor = FileProcessor()
