"""
Spreadsheet evidence as data, not prose.

A workbook is parsed ONCE into a :class:`TableSet`: typed cells, the real
spreadsheet row number of every row, a header and a kind per column. Everything
else is derived from that one parse, so nothing can disagree with it:

* :func:`redact_tableset` masks identifiers **by column**. The prose redactor
  decides from an 80-character context window, which in a table reaches the
  "Account No" header from the next three rows and masks phone numbers and
  amounts. Here each cell is judged with its own column header as its only
  context, so a column is masked consistently and its neighbours never are.
* :func:`render_text` produces the text that is chunked, embedded and put in
  prompts. It opens with a **workbook summary computed from every row** (row
  counts, per-column totals, date ranges, value counts). That summary is the
  first paragraph, so panel packing, the document classifier and chat all see
  exact whole-sheet figures even when they only read the opening of the
  document. Rows follow in blocks that each restate the sheet, the row range
  and the header row, so any chunk retrieved alone is still interpretable.
* :func:`row_records` / :func:`manifest` feed the row store, from which
  ``app/rag/table_query.py`` computes counts, sums and filters over every row
  deterministically. The model never adds up rows itself.

Assumptions, stated in the output where they matter:

* Formulas are read as the value Excel last saved. A formula with no saved
  value (a file written by a tool that never calculated it) is kept as its
  formula text, counted, and excluded from every figure.
* Dates typed as text are read day-first (DD/MM/YYYY), the Indian convention.
* A row whose first text cell starts with "Total" / "Grand total" /
  "Sub-total" is the sheet's own total. It stays in the text but is excluded
  from computed figures, which would otherwise count it twice.
"""

from __future__ import annotations

import datetime as _dt
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from loguru import logger

# Bounds on extraction. Every extracted character is chunked, embedded and
# stored, so an export with hundreds of thousands of rows would cost minutes
# of ingestion. Hitting a cap is stated in the text; nothing is cut silently.
MAX_CELLS = 200_000
MAX_CELL_CHARS = 2_000
# Rows are grouped into blocks of about this many characters. Kept under the
# chunker's TARGET_CHARS (app/rag/chunker.py, 1800) so a block becomes one
# chunk instead of being split mid-table.
BLOCK_CHARS = 1500
# The workbook summary is the document's first paragraph; this keeps it small
# enough to survive panel packing and the classifier's 6,000-character view.
SUMMARY_MAX_CHARS = 6000
# A text column with at most this many distinct values is summarised as value
# counts ("Paid 333, Pending 167") and offered to the query planner verbatim.
CATEGORY_MAX_DISTINCT = 12
# The header is looked for among this many leading non-empty rows.
HEADER_SCAN_ROWS = 10

KIND_NUMBER = "number"
KIND_DATE = "date"
KIND_BOOLEAN = "boolean"
KIND_TEXT = "text"
KIND_IDENTIFIER = "identifier"

#: Headers that name a measurement. Checked before the identifier test, so
#: "Invoice Amount" is a number and "Invoice No" is an identifier.
_MEASURE_HEADER_RE = re.compile(
    r"\b(amount|amt|value|total|sum|price|cost|fee|fees|salary|wage|balance|"
    r"qty|quantity|score|marks?|age|rate|percent(age)?|weight|height|count|"
    r"duration|days|hours|interest|principal|emi|tax|gst amount|paid|due|"
    r"outstanding|debit|credit|deposit|withdrawal)\b|%|₹|\brs\b|\binr\b",
    re.I,
)
#: Headers that name an identifier: never summed or averaged.
_ID_HEADER_RE = re.compile(
    r"\b(id|code|ref|reference|mobile|phone|tel|telephone|contact|pin|pincode|"
    r"zip|account|a/c|acct|aadhaar|aadhar|uid|pan|gstin|ifsc|cin|txn|utr|"
    r"invoice|cheque|chq|serial|roll|registration|reg|folio|policy|imei|"
    r"vehicle|fir|case)\b"
    r"|\b(no|nos|num|number)\b\.?(?!\s*of\b)"
    r"|^\s*(s|sr|sl)\.?\s*no\.?\s*$|^#$",
    re.I,
)
#: Cells meaning "no value" that must not turn a number column into text:
#: three "N/A" cells in a short column of amounts are still amounts.
_PLACEHOLDERS = frozenset(
    ["n/a", "na", "n.a.", "nil", "none", "-", "--", "—", "–", "not applicable",
     "not available", "tbd", "nan", "null", "#n/a", "#value!", "#ref!", "#div/0!"]
)
_TOTAL_ROW_RE = re.compile(r"^\s*(grand\s+|sub[\s-]?)?totals?\b", re.I)
_NUMBER_TEXT_RE = re.compile(
    r"^\(?\s*-?\s*(?:₹|rs\.?|inr)?\s*(\d{1,3}(?:,\d{2,3})+|\d+)(\.\d+)?\s*\)?\s*(%)?$",
    re.I,
)
_DMY_RE = re.compile(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2}|\d{4})$")
_ISO_RE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")
_MONTHS = {
    m: i
    for i, names in enumerate(
        [
            ("jan", "january"), ("feb", "february"), ("mar", "march"),
            ("apr", "april"), ("may",), ("jun", "june"), ("jul", "july"),
            ("aug", "august"), ("sep", "sept", "september"), ("oct", "october"),
            ("nov", "november"), ("dec", "december"),
        ],
        start=1,
    )
    for m in names
}
_DMONY_RE = re.compile(r"^(\d{1,2})(?:st|nd|rd|th)?[\s-]+([a-z]+)\.?,?[\s-]+(\d{4})$", re.I)


# ---------------------------------------------------------------------------
# Model
# ---------------------------------------------------------------------------


@dataclass
class Cell:
    """One cell: what the sheet shows, and its typed value for computation.

    ``value`` is JSON-safe: int/float for numbers, bool, an ISO string for
    dates ("2026-09-01" or "2026-09-01 10:30:00"), str for text, or None when
    the cell has no usable value (empty, or a formula with no saved value).
    """

    display: str
    value: Any = None
    kind: str = KIND_TEXT  # the cell's own type, before column inference
    formula: bool = False  # a formula whose value was never saved


@dataclass
class Column:
    name: str
    kind: str = KIND_TEXT
    letter: str = ""


@dataclass
class TableRow:
    number: int  # the spreadsheet's own row number (1-based)
    cells: List[Cell]
    is_total: bool = False


@dataclass
class Sheet:
    title: str
    index: int
    hidden: bool = False
    header_row: Optional[int] = None
    columns: List[Column] = field(default_factory=list)
    preamble: List[TableRow] = field(default_factory=list)  # title rows above the header
    rows: List[TableRow] = field(default_factory=list)
    uncomputed_formulas: int = 0
    day_first_dates: bool = False  # text dates were read as DD/MM/YYYY

    @property
    def data_rows(self) -> List[TableRow]:
        return [r for r in self.rows if not r.is_total]

    @property
    def total_rows(self) -> List[TableRow]:
        return [r for r in self.rows if r.is_total]


@dataclass
class TableSet:
    sheets: List[Sheet] = field(default_factory=list)
    cell_count: int = 0
    truncated: bool = False

    @property
    def is_empty(self) -> bool:
        return not any(s.rows or s.preamble for s in self.sheets)


# ---------------------------------------------------------------------------
# Cell typing
# ---------------------------------------------------------------------------


def _decimal(value: Any) -> Optional[Decimal]:
    """Exact decimal of a typed number; floats via their shortest repr, so the
    figure is exactly what the sheet displays (no 0.30000000000000004)."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None
        return Decimal(repr(value))
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def format_number(value: Any) -> str:
    """Plain rendering of a number: no grouping, no trailing zeros."""
    d = _decimal(value)
    if d is None:
        return ""
    if d == d.to_integral_value():
        return str(d.quantize(Decimal(1)))
    text = format(d.normalize(), "f")
    return text


def format_indian(value: Any, places: Optional[int] = None) -> str:
    """Indian digit grouping (12,34,56,789.50), as a lawyer would write it."""
    d = _decimal(value)
    if d is None:
        return ""
    if places is not None:
        d = d.quantize(Decimal(1).scaleb(-places))
    sign = "-" if d < 0 else ""
    text = format(abs(d), "f")
    whole, _, frac = text.partition(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        whole = ",".join(groups + [tail])
    return sign + whole + ("." + frac if frac else "")


def _clean_text(value: Any) -> str:
    text = re.sub(r"\s+", " ", str(value)).strip()
    if len(text) > MAX_CELL_CHARS:
        text = text[:MAX_CELL_CHARS] + "..."
    return text


def parse_date_text(text: str) -> Optional[str]:
    """ISO date from a typed date string, day-first; None if it is not one."""
    s = text.strip()
    m = _ISO_RE.match(s)
    if m:
        y, mo, d = (int(g) for g in m.groups())
    else:
        m = _DMY_RE.match(s)
        if m:
            d, mo, y = (int(g) for g in m.groups())
            if y < 100:
                y += 2000 if y < 70 else 1900
        else:
            m = _DMONY_RE.match(s)
            if not m:
                return None
            d, y = int(m.group(1)), int(m.group(3))
            mo = _MONTHS.get(m.group(2).lower())
            if mo is None:
                return None
    try:
        return _dt.date(y, mo, d).isoformat()
    except ValueError:
        return None


def parse_number_text(text: str) -> Optional[Tuple[Any, bool]]:
    """``(number, is_percent)`` from a number typed as text ("1,25,000",
    "₹ 500", "(1,200)", "15%"), or None. Strings with a leading zero ("007")
    are codes, not numbers."""
    s = text.strip()
    m = _NUMBER_TEXT_RE.match(s)
    if not m:
        return None
    digits = m.group(1).replace(",", "")
    if len(digits) > 1 and digits.startswith("0") and not m.group(2):
        return None
    number = Decimal(digits + (m.group(2) or ""))
    # Accounting negatives: "(1,200)" and "-1,200" / "₹ -500".
    if (s.startswith("(") and s.endswith(")")) or re.match(r"^\(?\s*-", s) or re.search(
        r"(?:₹|rs\.?|inr)\s*-", s, re.I
    ):
        number = -number
    value: Any = int(number) if number == number.to_integral_value() and not m.group(2) else float(number)
    return value, bool(m.group(3))


def make_cell(raw: Any, number_format: Optional[str] = None) -> Cell:
    """Type one openpyxl value, keeping its display form."""
    if raw is None:
        return Cell("", None, KIND_TEXT)
    if isinstance(raw, bool):
        return Cell("TRUE" if raw else "FALSE", raw, KIND_BOOLEAN)
    if isinstance(raw, _dt.datetime):
        if raw.time() == _dt.time(0):
            iso = raw.date().isoformat()
        else:
            iso = raw.isoformat(sep=" ", timespec="seconds")
        return Cell(iso, iso, KIND_DATE)
    if isinstance(raw, _dt.date):
        return Cell(raw.isoformat(), raw.isoformat(), KIND_DATE)
    if isinstance(raw, _dt.time):
        return Cell(raw.isoformat(), raw.isoformat(), KIND_TEXT)
    if isinstance(raw, _dt.timedelta):
        return Cell(str(raw), str(raw), KIND_TEXT)
    if isinstance(raw, (int, float)):
        if isinstance(raw, float) and (raw != raw or raw in (float("inf"), float("-inf"))):
            return Cell(str(raw), None, KIND_TEXT)
        value: Any = raw
        if isinstance(raw, float) and raw.is_integer() and abs(raw) < 1e15:
            value = int(raw)
        if number_format and "%" in number_format:
            # Stored as a fraction (0.155), shown as a percentage (15.5%).
            pct = _decimal(value) * 100
            value = int(pct) if pct == pct.to_integral_value() else float(pct)
            return Cell(format_number(value) + "%", value, KIND_NUMBER)
        return Cell(format_number(value), value, KIND_NUMBER)
    text = _clean_text(raw)
    if not text:
        return Cell("", None, KIND_TEXT)
    as_number = parse_number_text(text)
    if as_number is not None:
        return Cell(text, as_number[0], KIND_NUMBER)
    as_date = parse_date_text(text)
    if as_date is not None:
        return Cell(text, as_date, KIND_DATE)
    return Cell(text, text, KIND_TEXT)


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def _column_letter(index: int) -> str:
    out = ""
    n = index + 1
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out = chr(ord("A") + rem) + out
    return out


def _find_header(rows: Sequence[TableRow]) -> Optional[int]:
    """Index (into ``rows``) of the header row, or None for a headerless sheet.

    The header is the first of the leading rows that is as wide as the table
    (at least 60% of the widest early row, and 2+ cells) and is mostly text.
    Rows above it are titles ("Statement of account for FY 2025-26").
    """
    if not rows:
        return None
    window = rows[:30]
    widest = max(sum(1 for c in r.cells if c.display) for r in window)
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        filled = [c for c in row.cells if c.display]
        if len(filled) < max(2, int(widest * 0.6 + 0.999)):
            continue
        texty = sum(1 for c in filled if c.kind == KIND_TEXT)
        if texty >= max(1, int(len(filled) * 0.6 + 0.999)):
            return i
        return None  # the table starts with data: no header
    return None


def _infer_kind(name: str, cells: Iterable[Cell]) -> str:
    kinds: Dict[str, int] = {}
    filled = 0
    for c in cells:
        if c.value is None or (c.kind == KIND_TEXT and c.display.casefold() in _PLACEHOLDERS):
            continue
        filled += 1
        kinds[c.kind] = kinds.get(c.kind, 0) + 1
    measure = bool(_MEASURE_HEADER_RE.search(name))
    if not measure and _ID_HEADER_RE.search(name):
        return KIND_IDENTIFIER
    if not filled:
        return KIND_TEXT
    kind, count = max(kinds.items(), key=lambda kv: kv[1])
    if count / filled >= 0.8:
        return kind
    return KIND_TEXT


def _finish_sheet(sheet: Sheet, raw_rows: List[TableRow], lead: int) -> None:
    header_at = _find_header(raw_rows)
    if header_at is not None and header_at == len(raw_rows) - 1:
        # Nothing below it: the only row is data, not a header.
        header_at = None
    width = max((len(r.cells) for r in raw_rows), default=0)
    if header_at is not None:
        header_row = raw_rows[header_at]
        sheet.header_row = header_row.number
        sheet.preamble = raw_rows[:header_at]
        body = raw_rows[header_at + 1:]
        names = [c.display for c in header_row.cells]
    else:
        body = raw_rows
        names = []
    seen: Dict[str, int] = {}
    columns: List[Column] = []
    for i in range(width):
        letter = _column_letter(i + lead)
        name = names[i] if i < len(names) and names[i] else f"Column {letter}"
        key = name.lower()
        if key in seen:
            seen[key] += 1
            name = f"{name} ({seen[key]})"
        else:
            seen[key] = 1
        columns.append(Column(name=name, letter=letter))
    for row in body:
        row.cells.extend(Cell("") for _ in range(width - len(row.cells)))
        first_text = next((c.display for c in row.cells if c.display), "")
        row.is_total = bool(_TOTAL_ROW_RE.match(first_text)) and any(
            c.kind == KIND_NUMBER for c in row.cells
        )
    data = [r for r in body if not r.is_total]
    for i, col in enumerate(columns):
        col.kind = _infer_kind(col.name, (r.cells[i] for r in data))
        if col.kind == KIND_DATE:
            sheet.day_first_dates = sheet.day_first_dates or any(
                r.cells[i].kind == KIND_DATE and r.cells[i].display != r.cells[i].value
                for r in data
            )
    sheet.columns = columns
    sheet.rows = body


def parse_workbook(data: bytes) -> TableSet:
    """Parse an XLSX/XLSM workbook into typed sheets. Raises on an unreadable file.

    Two read-only passes run in lockstep: values (``data_only=True``) for what
    the sheet shows, formulas for spotting a formula whose value was never
    saved, which would otherwise read as an empty cell and silently drop out of
    every total.
    """
    import openpyxl

    values_wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    formulas_wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False)
    out = TableSet()
    try:
        for index, (ws, fws) in enumerate(zip(values_wb.worksheets, formulas_wb.worksheets)):
            # Some generators write a wrong <dimension>; without a reset,
            # read-only mode trusts it and silently skips rows past it.
            ws.reset_dimensions()
            fws.reset_dimensions()
            sheet = Sheet(title=str(ws.title), index=index, hidden=ws.sheet_state != "visible")
            raw_rows: List[TableRow] = []
            formula_rows = fws.iter_rows(values_only=True)
            for number, row in enumerate(ws.iter_rows(), start=1):
                frow = next(formula_rows, ())
                cells: List[Cell] = []
                for j, cell in enumerate(row):
                    raw = getattr(cell, "value", None)
                    fmt = getattr(cell, "number_format", None)
                    made = make_cell(raw, fmt)
                    if raw is None and j < len(frow):
                        formula = frow[j]
                        if isinstance(formula, str) and formula.startswith("="):
                            made = Cell(_clean_text(formula), None, KIND_TEXT, formula=True)
                    cells.append(made)
                while cells and not cells[-1].display:
                    cells.pop()
                if not cells:
                    continue
                filled = sum(1 for c in cells if c.display)
                if out.cell_count + filled > MAX_CELLS:
                    out.truncated = True
                    break
                out.cell_count += filled
                sheet.uncomputed_formulas += sum(1 for c in cells if c.formula)
                raw_rows.append(TableRow(number=number, cells=cells))
            if raw_rows:
                # Drop columns empty in every row (a table that starts at C).
                lead = min(next(i for i, c in enumerate(r.cells) if c.display) for r in raw_rows)
                for r in raw_rows:
                    r.cells = r.cells[lead:]
                _finish_sheet(sheet, raw_rows, lead)
            out.sheets.append(sheet)
            if out.truncated:
                break
    finally:
        values_wb.close()
        formulas_wb.close()
    return out


# ---------------------------------------------------------------------------
# Redaction (column-aware)
# ---------------------------------------------------------------------------


def _redact_cell(cell: Cell, header: str, matches: list) -> Cell:
    from app.processing.pii_detector import scan_text

    text = cell.display
    # Every identifier is 8+ characters and contains a digit.
    if len(text) < 8 or not any(ch.isdigit() for ch in text):
        return cell
    prefix = f"{header}: " if header else ""
    scan = scan_text(prefix + text, strict_context=True)
    hits = [m for m in scan.matches if m.span[0] >= len(prefix)]
    if not hits:
        return cell
    masked, cursor = [], 0
    for m in sorted(hits, key=lambda m: (m.span[0], -(m.span[1] - m.span[0]))):
        start, end = m.span[0] - len(prefix), m.span[1] - len(prefix)
        if start < cursor:
            continue
        masked.append(text[cursor:start])
        masked.append(m.masked)
        cursor = end
    masked.append(text[cursor:])
    matches.extend(hits)
    redacted = "".join(masked)
    # A masked value is no longer a number or a date: it stays a string.
    return Cell(redacted, redacted, KIND_TEXT)


def redact_tableset(tables: TableSet):
    """A masked copy of ``tables`` plus the identifier matches found.

    Each cell is scanned with only its own column header as context (see
    ``scan_text(strict_context=True)``), so "Account No" masks its column
    consistently and never reaches the Mobile or Amount column beside it.
    Returns ``(TableSet, PIIScanResult)``; the scan result's ``redacted_text``
    is filled in by the caller from :func:`render_text`.
    """
    from app.processing.pii_detector import PIIScanResult

    matches: list = []
    sheets: List[Sheet] = []
    for s in tables.sheets:
        headers = [c.name for c in s.columns]

        def redact_rows(rows: List[TableRow], with_headers: bool) -> List[TableRow]:
            out = []
            for r in rows:
                cells = [
                    _redact_cell(c, headers[i] if with_headers and i < len(headers) else "", matches)
                    for i, c in enumerate(r.cells)
                ]
                out.append(TableRow(number=r.number, cells=cells, is_total=r.is_total))
            return out

        # Header names are scanned too: "Payments to A/c 1234567890" is a header.
        columns = [
            Column(name=_redact_cell(Cell(c.name, c.name), "", matches).display, kind=c.kind, letter=c.letter)
            for c in s.columns
        ]
        sheets.append(
            Sheet(
                title=s.title,
                index=s.index,
                hidden=s.hidden,
                header_row=s.header_row,
                columns=columns,
                preamble=redact_rows(s.preamble, with_headers=False),
                rows=redact_rows(s.rows, with_headers=True),
                uncomputed_formulas=s.uncomputed_formulas,
                day_first_dates=s.day_first_dates,
            )
        )
    return (
        TableSet(sheets=sheets, cell_count=tables.cell_count, truncated=tables.truncated),
        PIIScanResult(redacted_text="", matches=matches),
    )


# ---------------------------------------------------------------------------
# Column statistics (shared by the summary and the query engine)
# ---------------------------------------------------------------------------


def column_values(sheet: Sheet, index: int) -> List[Cell]:
    return [r.cells[index] for r in sheet.data_rows if index < len(r.cells)]


def describe_column(sheet: Sheet, index: int) -> str:
    """One summary line for a column, computed from every data row."""
    col = sheet.columns[index]
    cells = column_values(sheet, index)
    filled = [c for c in cells if c.display]
    parts = [f"{len(filled)} filled"]
    if col.kind == KIND_NUMBER:
        nums = [d for d in (_decimal(c.value) for c in filled if c.kind == KIND_NUMBER) if d is not None]
        if nums:
            parts.append(f"sum {format_number(sum(nums))}")
            parts.append(f"min {format_number(min(nums))}")
            parts.append(f"max {format_number(max(nums))}")
        skipped = len(filled) - len(nums)
        if skipped:
            parts.append(f"{skipped} non-numeric cell(s) not counted")
    elif col.kind == KIND_DATE:
        dates = sorted(str(c.value)[:10] for c in filled if c.kind == KIND_DATE and c.value)
        if dates:
            parts.append(f"{dates[0]} to {dates[-1]}")
    elif col.kind in (KIND_TEXT, KIND_BOOLEAN):
        counts: Dict[str, int] = {}
        for c in filled:
            counts[c.display] = counts.get(c.display, 0) + 1
        if 0 < len(counts) <= CATEGORY_MAX_DISTINCT:
            ranked = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))
            parts.append(", ".join(f"{k} {v}" for k, v in ranked))
        else:
            parts.append(f"{len(counts)} distinct")
    elif col.kind == KIND_IDENTIFIER:
        parts.append(f"{len({c.display for c in filled})} distinct")
    return f"  {col.name} ({col.kind}): " + "; ".join(parts)


def _sheet_heading(sheet: Sheet) -> str:
    data = sheet.data_rows
    name = f'Sheet "{sheet.title}"' + (" (hidden)" if sheet.hidden else "")
    if not data:
        return f"{name}: no data rows"
    first, last = data[0].number, data[-1].number
    span = f"row {first}" if first == last else f"rows {first}-{last}"
    line = f"{name}: {len(data)} data row(s) ({span}), {len(sheet.columns)} column(s)"
    totals = sheet.total_rows
    if totals:
        line += (
            f"; {len(totals)} total row(s) ("
            + ", ".join(f"row {r.number}" for r in totals)
            + ") excluded from the figures below"
        )
    if sheet.uncomputed_formulas:
        line += (
            f"; {sheet.uncomputed_formulas} formula cell(s) have no saved value and "
            "are excluded"
        )
    if sheet.day_first_dates:
        line += "; dates written as text were read as DD/MM/YYYY"
    return line


def summary_text(tables: TableSet) -> str:
    """The workbook summary paragraph: exact figures over every data row."""
    lines = ["Workbook summary (computed by the system from every row; exact)"]
    for sheet in tables.sheets:
        if not (sheet.rows or sheet.preamble):
            continue
        lines.append(_sheet_heading(sheet))
        if sheet.data_rows:
            for i in range(len(sheet.columns)):
                lines.append(describe_column(sheet, i))
    if tables.truncated:
        lines.append(
            f"Extraction stopped after {MAX_CELLS} cells; rows after that point "
            "were not read and are not in these figures."
        )
    text = "\n".join(lines)
    if len(text) > SUMMARY_MAX_CHARS:
        cut = text.rfind("\n", 0, SUMMARY_MAX_CHARS - 80)
        text = text[: cut if cut > 0 else SUMMARY_MAX_CHARS - 80] + (
            "\n  ... further columns and sheets omitted from this summary; every "
            "row is still stored and searchable"
        )
    return text


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


@dataclass
class RenderedRow:
    """Where one row's line sits in the rendered text."""

    offset: int
    sheet: str
    row: int


def render_text(tables: TableSet) -> Tuple[str, List[RenderedRow]]:
    """Text for chunking, embedding and prompts, plus the offset of every row
    line in it (so a citation can name the exact spreadsheet row)."""
    parts: List[str] = []
    index: List[RenderedRow] = []
    pos = 0

    def add(text: str, rows: List[Tuple[int, int]] = ()) -> None:
        nonlocal pos
        if parts:
            pos += 2  # the "\n\n" separator
        for rel, number in rows:
            index.append(RenderedRow(pos + rel, current_sheet, number))
        parts.append(text)
        pos += len(text)

    current_sheet = ""
    add(summary_text(tables))
    for sheet in tables.sheets:
        current_sheet = sheet.title
        label = f"Sheet: {sheet.title}" + (" (hidden)" if sheet.hidden else "")
        header_line = " | ".join(c.name for c in sheet.columns) if sheet.header_row else ""

        def lines_of(rows: List[TableRow]) -> List[Tuple[int, str]]:
            out = []
            for r in rows:
                cells = [c.display for c in r.cells]
                while cells and not cells[-1]:
                    cells.pop()
                out.append((r.number, " | ".join(cells)))
            return out

        groups: List[Tuple[bool, List[Tuple[int, str]]]] = []
        if sheet.preamble:
            groups.append((False, lines_of(sheet.preamble)))
        body = lines_of(sheet.rows)
        block: List[Tuple[int, str]] = []
        size = 0
        for number, line in body:
            if block and size + len(line) > BLOCK_CHARS:
                groups.append((True, block))
                block, size = [], 0
            block.append((number, line))
            size += len(line) + 1
        if block:
            groups.append((True, block))

        for with_header, rows in groups:
            first, last = rows[0][0], rows[-1][0]
            span = f"row {first}" if first == last else f"rows {first}-{last}"
            text_lines = [f"{label} ({span})"]
            if with_header and header_line:
                text_lines.append(header_line)
            offsets = []
            rel = sum(len(t) + 1 for t in text_lines)
            for number, line in rows:
                offsets.append((rel, number))
                text_lines.append(line)
                rel += len(line) + 1
            add("\n".join(text_lines), offsets)
    return "\n\n".join(parts), index


# ---------------------------------------------------------------------------
# Row store records
# ---------------------------------------------------------------------------


def manifest(tables: TableSet) -> Dict[str, Any]:
    """What the query planner may see about the workbook: sheets, columns,
    kinds, row ranges and, for categorical columns, the exact values."""
    sheets = []
    for s in tables.sheets:
        data = s.data_rows
        if not data:
            continue
        columns = []
        for i, col in enumerate(s.columns):
            entry: Dict[str, Any] = {"name": col.name, "kind": col.kind}
            if col.kind in (KIND_TEXT, KIND_BOOLEAN):
                distinct = sorted({c.display for c in column_values(s, i) if c.display})
                if len(distinct) <= 50:
                    entry["values"] = distinct[:50]
                else:
                    entry["examples"] = distinct[:5]
            columns.append(entry)
        sheets.append(
            {
                "sheet": s.title,
                "sheet_index": s.index,
                "hidden": s.hidden,
                "header_row": s.header_row,
                "first_row": data[0].number,
                "last_row": data[-1].number,
                "data_rows": len(data),
                "total_rows": [r.number for r in s.total_rows],
                "uncomputed_formulas": s.uncomputed_formulas,
                "day_first_dates": s.day_first_dates,
                "columns": columns,
            }
        )
    return {"sheets": sheets, "truncated": tables.truncated, "cell_count": tables.cell_count}


def row_records(tables: TableSet) -> List[Dict[str, Any]]:
    """One JSON-safe record per data row (total rows included, flagged)."""
    out = []
    for s in tables.sheets:
        for r in s.rows:
            out.append(
                {
                    "sheet": s.title,
                    "sheet_index": s.index,
                    "row": r.number,
                    "is_total": r.is_total,
                    "display": [c.display for c in r.cells],
                    "values": [c.value for c in r.cells],
                    "kinds": [c.kind for c in r.cells],
                }
            )
    return out


def extract_text(data: bytes) -> str:
    """Rendered text of a workbook (unredacted), for text-only callers."""
    tables = parse_workbook(data)
    text, _ = render_text(tables)
    logger.success(
        "Extracted {} characters from XLSX ({} sheets, {} cells)",
        len(text),
        len(tables.sheets),
        tables.cell_count,
    )
    return text
