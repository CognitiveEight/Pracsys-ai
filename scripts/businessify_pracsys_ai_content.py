#!/usr/bin/env python3
"""Scrub AI implementation detail from Pracsys_workflow_overview.docx.

The file is going to an external client. This script replaces ~50
technical mentions (model names, library names, internal event
identifiers, pipeline-node names, protocol details) with business
language so a reader cannot clone the system from the document.

Specifically:

* Drops §4.3.4 Socket Events entirely (literal API contract).
* Rewrites §4.3 high-level bullets, §4.3.1, §4.3.2, §4.3.3, §4.3.5,
  §5.3 Step 3/4, §5.4, §5.5, §6.2, §7.2 paragraphs in place.
* Renames §4.3.2 → "AI Capabilities" and §4.3.3 → "AI Deliverables".
* Removes the five AI-tech glossary entries (Gemini, Qdrant, LangGraph,
  Vertex AI Search, Chirp 3).
* Tweaks two §1 summary bullets.

Idempotent: detects the post-businessify state on entry and exits.

Run from the repo root:

    python3 scripts/businessify_pracsys_ai_content.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(__file__).resolve().parent.parent / "Pracsys_workflow_overview.docx"


def _normalize(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().lower()


def _find(doc, prefix: str, *, optional: bool = False):
    needle = _normalize(prefix)
    for p in doc.paragraphs:
        if _normalize(p.text).startswith(needle):
            return p
    if optional:
        return None
    raise LookupError(f"not found: {prefix!r}")


def _set_run_text(run_elem, text: str) -> None:
    for t in list(run_elem.findall(qn("w:t"))):
        run_elem.remove(t)
    new_t = OxmlElement("w:t")
    new_t.text = text
    new_t.set(qn("xml:space"), "preserve")
    run_elem.append(new_t)


def replace_text(p: Paragraph, text: str) -> None:
    """Rewrite a paragraph's entire text into its first run."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def replace_two_run_bullet(p: Paragraph, label: str, rest: str) -> None:
    """Rewrite a two-run "Label: detail" bullet (bold prefix + regular rest)."""
    runs = p._p.findall(qn("w:r"))
    if len(runs) >= 2:
        _set_run_text(runs[0], f"{label}: ")
        _set_run_text(runs[1], rest)
        for r in runs[2:]:
            for t in list(r.findall(qn("w:t"))):
                r.remove(t)
    else:
        replace_text(p, f"{label}: {rest}")


def delete_paragraph(p: Paragraph) -> None:
    p._p.getparent().remove(p._p)


def delete_from_inclusive_to_exclusive(start_p: Paragraph, end_p: Paragraph) -> int:
    body = start_p._p.getparent()
    cur = start_p._p
    removed = 0
    while cur is not None and cur is not end_p._p:
        nxt = cur.getnext()
        body.remove(cur)
        cur = nxt
        removed += 1
    return removed


def main() -> int:
    if not DOCX_PATH.exists():
        print(f"ERROR: missing {DOCX_PATH}", file=sys.stderr)
        return 1

    doc = Document(str(DOCX_PATH))

    # Idempotency: if §4.3.4 is gone and §4.3.3 has been renamed, we're done.
    if (
        _find(doc, "4.3.4 Socket Events", optional=True) is None
        and _find(doc, "4.3.3 AI Deliverables", optional=True) is not None
    ):
        print("already business-ified — no changes")
        return 0

    # Sanity check pre-conditions from the previous step.
    try:
        _find(doc, "4.3.5 Indian-Law Grounding")
        _find(doc, "4.3.1 Technology Stack")
        _find(doc, "4.3.4 Socket Events")
    except LookupError as exc:
        print(f"docx not in expected state: {exc}", file=sys.stderr)
        return 2

    # ===== §1 summary tweaks ===============================================
    replace_text(
        _find(doc, "Strategy & Objective"),
        "Strategy & Objective — three to five case-tailored strategies or opinions per workspace, refined as more evidence is uploaded.",
    )
    replace_text(
        _find(doc, "Drafting Assistant — AI-suggested document drafts"),
        "Drafting Assistant — AI-suggested document drafts (notices, applications, replies) rendered into the lawyer's own DOCX template, with logo, fonts, headers and statutory boilerplate preserved exactly.",
    )

    # ===== §4.3 high-level bullets ========================================
    replace_two_run_bullet(
        _find(doc, "Runtime: Independent Python"),
        "Runtime",
        "A dedicated, isolated AI service deployed in our cloud environment; it starts independently of the Backend and connects to it over a real-time channel.",
    )
    replace_two_run_bullet(
        _find(doc, "Channel: Socket.IO over an authenticated"),
        "Channel",
        "An authenticated, server-to-server real-time channel between the Backend and the AI service (authentication detail in §7.2).",
    )
    replace_two_run_bullet(
        _find(doc, "AI Stack: Google Gemini"),
        "AI Stack",
        "Built on enterprise-grade managed AI services from a leading cloud provider, combined with proprietary legal-reasoning components developed by SimplerToday. No customer data is used to train shared models.",
    )
    # Boundary and Responsibilities already generic — leave them alone.

    # ===== §4.3.1 Technology Stack — 12 bullets ============================
    s4_3_1 = [
        ("Language & Runtime", "Hosting", "Runs on isolated cloud infrastructure, separate from the public API tier and the Backend datastore."),
        ("LLM", "Reasoning & Drafting AI", "An enterprise large language model from a leading cloud provider; used for case reasoning, classification, drafting, and audio refinement."),
        ("Embeddings", "Document Understanding", "Enterprise-grade document embeddings used for semantic search and retrieval."),
        ("Vector Store", "Knowledge Index", "A secure vector-based document retrieval layer, strictly scoped per workspace so workspaces remain isolated from one another."),
        ("Agent Orchestration", "AI Orchestration", "A proprietary multi-step legal-reasoning orchestrator developed by SimplerToday."),
        ("Law-Citation Retrieval", "Indian-Law Citation Sources", "A leading public Indian case-law and statutes repository, plus a curated commercial legal-search service operated by a major cloud provider."),
        ("Speech-to-Text", "Audio Transcription", "An enterprise speech recognition service with built-in speaker identification and broad Indian-language coverage."),
        ("Audio Pre-processing", "Audio Pre-processing", "Industry-standard audio format normalisation."),
        ("Document Parsing", "Document Parsing & OCR", "Industry-standard parsers and OCR across PDF, DOCX, image, and scanned formats."),
        ("Template Rendering", "Drafting Engine", "A proprietary template-preservation engine that renders into the lawyer's own DOCX template without altering its formatting."),
        ("Transport", "Real-Time Transport", "An encrypted bidirectional channel to the Backend."),
        ("Logging & Observability", "Observability", "Centralised log shipping with request-ID correlation across services."),
    ]
    for old_label, new_label, new_rest in s4_3_1:
        replace_two_run_bullet(_find(doc, f"{old_label}:"), new_label, new_rest)

    # ===== §4.3.2 Agents and Orchestration → "AI Capabilities" =============
    replace_text(
        _find(doc, "4.3.2 Agents and Orchestration"),
        "4.3.2 AI Capabilities",
    )
    s4_3_2 = [
        ("Role Agent", "Case-Role Awareness", "Each workspace records the perspective the firm represents in the matter. All AI output is generated with respect to that perspective; changing it invalidates earlier output so it is regenerated cleanly rather than mixed across perspectives."),
        ("Intelligence Chat Agent", "Conversational Assistance", "A guarded conversational interface that answers lawyer questions, cites the underlying record, and routes drafting / validation / strategy / insight-removal requests to the appropriate specialised capability. Replies are deliberately neutral and accuracy-first; perspective-conditioned advocacy lives in the panel outputs (strategy, drafts), not in chat."),
        ("Document Assessor", "Document Quality Assessment", "Per-document readability scoring, quality-issue flagging, and evidentiary-strength assessment."),
        ("Workspace Metadata Generator", "Workspace Summarisation", "Automatic derivation of the case title, jurisdiction, category, and short description from the document corpus."),
        ("Strategy & Objective Generator", "Strategy Generation", "Case-tailored strategies and objectives per workspace, refined as more evidence is uploaded."),
        ("Analysis Report Generator", "Analytical Report", "Quality-scored case analysis with strengths, weaknesses, and missing-information leads."),
        ("List-of-Dates Generator", "Chronological Timeline", "Source-cited event timeline extracted from the documents."),
        ("Applicable-Laws Pipeline", "Applicable-Laws Extraction", "Indian-law citations and case-law mapping (details in §4.3.4)."),
        ("Procedural Timeline Generator", "Procedural Calendar", "Inferred upcoming filing, hearing, and compliance dates."),
        ("Drafting Agent", "Drafting Capability", "AI-fillable drafts rendered into the lawyer's own DOCX template, preserving formatting exactly (details in §5.5)."),
        ("Audio Transcription Agent", "Audio Evidence Transcription", "Automated transcription with speaker identification and content-aware refinement for multi-speaker calls."),
    ]
    for old_label, new_label, new_rest in s4_3_2:
        replace_two_run_bullet(_find(doc, f"{old_label}:"), new_label, new_rest)

    # ===== §4.3.3 Feature Pipeline and Outputs → "AI Deliverables" =========
    replace_text(
        _find(doc, "4.3.3 Feature Pipeline and Outputs"),
        "4.3.3 AI Deliverables",
    )
    # Five existing bullets → three. Rewrite the first three; delete the last two.
    replace_two_run_bullet(
        _find(doc, "Trigger Events:"),
        "Deliverables",
        "Each workspace receives eight AI deliverables: Document Assessments, Workspace Summary, Strategy & Objective, Analytical Report, List of Dates, Applicable Laws, Procedural Timeline, and Drafting Suggestions. Each is delivered to the Backend as it is produced.",
    )
    replace_two_run_bullet(
        _find(doc, "featureType Values:"),
        "Triggers",
        "AI processing begins when a workspace is created or when its document set changes. The lawyer is never blocked waiting for analyses to complete.",
    )
    replace_two_run_bullet(
        _find(doc, "Dispatch Model:"),
        "Resilience",
        "Long-running analyses run asynchronously and do not block conversational responses. AI deliverables are durable — if the connection to the Backend drops mid-delivery, the AI service re-delivers the latest outputs on reconnect, and they survive a full service restart without re-running the pipeline.",
    )
    delete_paragraph(_find(doc, "Replay Cache:"))
    delete_paragraph(_find(doc, "Persistent Snapshots:"))

    # ===== Drop §4.3.4 Socket Events entirely ==============================
    sec_4_3_4 = _find(doc, "4.3.4 Socket Events")
    sec_4_3_5 = _find(doc, "4.3.5 Indian-Law Grounding")
    removed_4_3_4 = delete_from_inclusive_to_exclusive(sec_4_3_4, sec_4_3_5)

    # Renumber §4.3.5 to §4.3.4 now that the old §4.3.4 is gone.
    replace_text(
        sec_4_3_5,
        "4.3.4 Indian-Law Grounding and Applicable-Laws Sourcing",
    )

    # ===== §4.3.5 (now §4.3.4) Indian-Law Grounding — light scrub ==========
    replace_text(
        _find(doc, "Applicable-laws output is grounded"),
        "Applicable-laws output is grounded against live Indian-law sources rather than relying on the AI's pre-training alone. Two retrievers feed the AI before the citation list is composed:",
    )
    replace_text(
        _find(doc, "Indian Kanoon"),
        "Primary source — a leading public Indian case-law and statutes repository, queried with a site-restricted lookup.",
    )
    replace_text(
        _find(doc, "Google Vertex AI Search"),
        "Supplementary source — a curated commercial legal-search service operated by a major cloud provider, configured against a tailored Indian-law data store.",
    )
    # "Two hard constraints govern the final emitted list:" — keep as is.
    replace_text(
        _find(doc, "Indian Penal Code, Code of Criminal Procedure"),
        "Indian Penal Code, Code of Criminal Procedure, and Indian Evidence Act are filtered out post-processing so they never reach the lawyer. The Bharatiya Nyaya Sanhita (BNS), Bharatiya Nagarik Suraksha Sanhita (BNSS), and Bharatiya Sakshya Adhiniyam (BSA) are the operative statutes for post-1 July 2024 matters and are the only criminal-law framing allowed in output.",
    )
    replace_text(
        _find(doc, "The applicable-laws list is capped"),
        "The applicable-laws list is capped at twelve items with stable per-item identifiers preserved across re-runs, so the lawyer's highlights and notes on a given law slot are not disturbed when the list is refreshed.",
    )

    # ===== §5.3 Step 3 — AI bootstrapping ==================================
    replace_text(
        _find(doc, "Step 3 — AI bootstrapping"),
        "Step 3 — AI bootstrapping. The Backend signals the AI service that a new workspace exists and shares short-lived signed URLs for the evidence files. The AI service first validates and locks the workspace's case role (the perspective the firm represents in the matter) before any analytical capability runs; subsequent role changes invalidate previously-generated outputs so they are regenerated cleanly rather than mixed across perspectives. The user does not need to wait — processing is asynchronous and the lawyer can begin interacting with the workspace immediately.",
    )

    # ===== §5.3 Step 4 — Progressive enrichment ============================
    replace_text(
        _find(doc, "Step 4 — Progressive enrichment"),
        "Step 4 — Progressive enrichment. As each AI capability produces its deliverable, the result is sent back to the Backend and merged into the workspace using domain-specific reconciliation rules. Long-running analyses run asynchronously without blocking conversational responses. AI deliverables survive transient disconnections — the AI service re-delivers the latest outputs on reconnect so the client view stays consistent. The Backend:",
    )

    # ===== §5.4 Intelligence Chat ==========================================
    replace_text(
        _find(doc, "Inside any workspace the lawyer can ask"),
        "Inside any workspace the lawyer can ask natural-language questions (\"Summarise the prosecution's argument\", \"Draft a reply to the notice dated 12 March\"). The AI assembles the relevant case context (case role, document excerpts, active strategy, approved facts, latest analysis, and any inline attachment) and produces an answer through a guarded multi-step pipeline gated by the same case-role lock as the rest of the AI service. The flow:",
    )
    replace_text(
        _find(doc, "The Backend persists the message"),
        "The Backend persists the message, then forwards it to the AI service. The AI service assembles the relevant case context (case role, document excerpts, active strategy, approved facts, latest analysis, any inline attachment) and runs it through its conversational pipeline.",
    )
    replace_text(
        _find(doc, "Every chat turn is durably stored"),
        "Every chat turn is durably stored against the workspace and contributes to the workspace's audit history. Chat replies are deliberately neutral and accuracy-first — perspective-conditioned advocacy lives in the panel outputs (strategy, drafts), not in the chat answer. State-mutating requests (drafting, validation, strategy change, insight removal) are routed to the appropriate specialised capability instead of being answered inline, and a governance review step ensures these routings are honoured.",
    )

    # ===== §5.5 Drafting Assistant =========================================
    replace_text(
        _find(doc, "The Drafting Assistant lets a lawyer convert"),
        "The Drafting Assistant lets a lawyer convert any of their own DOCX templates (notices, replies, applications) into a reusable, AI-fillable draft without compromising the firm's letterhead, fonts, logo, headers/footers, or statutory boilerplate.",
    )
    replace_text(
        _find(doc, "Step 1 — Template ingestion"),
        "Step 1 — Template ingestion (one-time per template). The AI analyses the lawyer's DOCX template and identifies the regions that vary between drafts: recipient name, case number, dates, factual narrative, prayer wording, signatures, amounts. Template-fixed content (firm letterhead, advocate roll number, statutory citation text, sign-offs) is preserved untouched.",
    )
    replace_text(
        _find(doc, "Step 2 — In-place"),
        "Step 2 — In-place annotation. The variable regions are annotated inside the template without altering surrounding formatting. Font, size, weight, colour, spacing, headers/footers, and any logo are preserved exactly.",
    )
    replace_text(
        _find(doc, "Step 3 — Per-draft fill"),
        "Step 3 — Per-draft fill. For each new draft, the AI fills the variable regions against the case facts (Indian-law grounded, named parties, formatted dates) and a proprietary template-preservation engine renders the final document. The output is visually identical to the lawyer's original template except for the variable text.",
    )
    replace_text(
        _find(doc, "Step 4 — Storage"),
        "Step 4 — Storage and delivery. The rendered DOCX is uploaded to Google Cloud Storage and a short-lived pre-signed URL is returned to the frontend, identical to the evidence-file access pattern described in §6.2.",
    )

    # ===== §6.2 Audio Files ================================================
    replace_text(
        _find(doc, "Audio Files:"),
        "Audio Files: Audio uploads (mp3, wav, m4a, ogg, flac, opus, webm, aac) are transcribed by the AI service using an enterprise-grade speech recognition service with built-in speaker identification and broad Indian-language support. For multi-speaker legal calls, a content-aware refinement step improves speaker attribution. The transcript is rendered as a timestamped, speaker-tagged narrative and becomes the document's text equivalent — so quality scoring, search, and chat all work on audio evidence without special-casing. A transcription-status indicator (ready / pending / failed) is shown per audio entry. Both the source audio and the generated transcript are stored, with explicit linkage via the source file's hash so provenance is preserved.",
    )

    # ===== §7.2 Authentication =============================================
    replace_text(
        _find(doc, "Authentication: The Socket.IO handshake"),
        "Authentication: The Backend-to-AI handshake is authenticated using a cryptographic signature-based credential exchange. The AI service is not exposed to the public internet; only the Backend is whitelisted to connect.",
    )

    # ===== §11 Glossary — remove five AI-tech entries ======================
    glossary_to_remove = [
        "Google Gemini —",
        "Qdrant —",
        "LangGraph —",
        "Vertex AI Search —",
        "Chirp 3 —",
    ]
    removed_glossary = 0
    for prefix in glossary_to_remove:
        p = _find(doc, prefix, optional=True)
        if p is not None:
            delete_paragraph(p)
            removed_glossary += 1

    doc.save(str(DOCX_PATH))
    print(f"Business-ified AI content in {DOCX_PATH.name}")
    print("  rewrote §1 summary (2 tweaks)")
    print("  rewrote §4.3 high-level bullets (3 of 5)")
    print("  rewrote §4.3.1 Technology Stack (12 bullets)")
    print("  renamed §4.3.2 → AI Capabilities, rewrote 11 bullets")
    print("  renamed §4.3.3 → AI Deliverables, compressed 5 → 3 bullets")
    print(f"  deleted §4.3.4 Socket Events ({removed_4_3_4} paragraphs)")
    print("  renumbered Indian-Law Grounding to §4.3.4 and lightly scrubbed it")
    print("  rewrote §5.3 Step 3, Step 4")
    print("  rewrote §5.4 Intelligence Chat (3 paragraphs)")
    print("  rewrote §5.5 Drafting Assistant (5 paragraphs)")
    print("  rewrote §6.2 Audio Files")
    print("  rewrote §7.2 Authentication")
    print(f"  removed {removed_glossary} glossary tech entries")
    return 0


if __name__ == "__main__":
    sys.exit(main())
