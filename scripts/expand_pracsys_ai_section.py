#!/usr/bin/env python3
"""Expand §4.3 of Pracsys_workflow_overview.docx into a full AI sub-chapter.

Inserts four new Heading 2 subsections under §4.3:
    4.3.1 Technology Stack
    4.3.2 Agents and Orchestration
    4.3.3 Feature Pipeline and Outputs
    4.3.4 Socket Events — Inbound and Outbound

Renames the existing §4.3.1 Applicable-Laws Sourcing heading to §4.3.5
Indian-Law Grounding and Applicable-Laws Sourcing so it becomes the last
subsection of the AI sub-chapter.

Idempotent: re-running after a successful run is a no-op.

Run from the repo root:

    python3 scripts/expand_pracsys_ai_section.py
"""

from __future__ import annotations

import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(__file__).resolve().parent.parent / "Pracsys_workflow_overview.docx"


def _normalize(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().lower()


def _find_paragraph_starting_with(doc, prefix: str) -> Paragraph:
    needle = _normalize(prefix)
    for para in doc.paragraphs:
        if _normalize(para.text).startswith(needle):
            return para
    raise LookupError(f"Could not locate paragraph starting with: {prefix!r}")


def _find_paragraph_starting_with_optional(doc, prefix: str):
    needle = _normalize(prefix)
    for para in doc.paragraphs:
        if _normalize(para.text).startswith(needle):
            return para
    return None


def _replace_paragraph_text(paragraph: Paragraph, new_text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def _insert_paragraph_after(
    anchor: Paragraph, text: str, *, style_source: Paragraph
) -> Paragraph:
    new_p = deepcopy(style_source._p)
    for r in list(new_p.findall(qn("w:r"))):
        new_p.remove(r)
    rpr_copy = None
    src_first_r = style_source._p.find(qn("w:r"))
    if src_first_r is not None:
        rpr = src_first_r.find(qn("w:rPr"))
        if rpr is not None:
            rpr_copy = deepcopy(rpr)
    new_r = OxmlElement("w:r")
    if rpr_copy is not None:
        new_r.append(rpr_copy)
    new_t = OxmlElement("w:t")
    new_t.text = text
    new_t.set(qn("xml:space"), "preserve")
    new_r.append(new_t)
    new_p.append(new_r)
    anchor._p.addnext(new_p)
    return Paragraph(new_p, anchor._parent)


def main() -> int:
    if not DOCX_PATH.exists():
        print(f"ERROR: missing {DOCX_PATH}", file=sys.stderr)
        return 1

    doc = Document(str(DOCX_PATH))

    # Idempotency — if §4.3.5 already exists, do nothing.
    if _find_paragraph_starting_with_optional(doc, "4.3.5 Indian-Law Grounding") is not None:
        print("already expanded — no changes")
        return 0

    # Sanity-check that the previous inline-update has already run.
    try:
        boundary = _find_paragraph_starting_with(doc, "Boundary: The AI service never")
        existing_4_3_1 = _find_paragraph_starting_with(doc, "4.3.1 Applicable-Laws Sourcing")
        _find_paragraph_starting_with(doc, "AI Stack: Google Gemini")
    except LookupError as exc:
        print(
            "ERROR: docx is not in the expected post-update state — "
            f"{exc}. Run scripts/update_pracsys_overview.py first.",
            file=sys.stderr,
        )
        return 2

    # Style templates.
    heading2_template = existing_4_3_1                 # known Heading 2 with 4.3.x numbering
    normal_template = _find_paragraph_starting_with(   # known Normal paragraph
        doc, "Applicable-laws output is grounded"
    )
    bullet_template = boundary                          # known List Bullet paragraph

    last = boundary  # all new content goes between `boundary` and `existing_4_3_1`

    # ===== §4.3.1 Technology Stack ==========================================
    last = _insert_paragraph_after(
        last, "4.3.1 Technology Stack", style_source=heading2_template
    )
    last = _insert_paragraph_after(
        last,
        "The AI service is built on the following enterprise-grade components. Every dependency is either a Google Cloud managed service or a mature open-source library; no in-house cryptography or model training is involved.",
        style_source=normal_template,
    )
    tech_stack_bullets = [
        "Language and runtime — Python 3.11+ asyncio service; single long-lived process per environment.",
        "LLM (reasoning, classification, drafting, audio re-diarization) — Google Gemini, accessed through the official Google generative AI SDK.",
        "Embeddings — Google text-embedding-004, 768-dimensional vectors, cosine similarity.",
        "Vector store — Qdrant; a single collection holds every workspace's document chunks and feature-state snapshots, with a workspace_id payload key that filters every read and write so workspaces are strictly isolated.",
        "Agent orchestration — LangGraph (state-machine framework) for the role agent and the intelligence-chat graph.",
        "Law-citation retrieval — Google Vertex AI Search (langchain_google_community.VertexAISearchRetriever against a curated data store) as primary; Indian Kanoon (indiankanoon.org) site-restricted retrieval as fallback.",
        "Speech-to-text — Google Speech-to-Text V2 (Chirp 3) with GCS-backed BatchRecognize for native speaker diarization.",
        "Audio pre-processing — ffmpeg (on PATH) for 16-bit PCM, 16 kHz mono normalisation. pyannote.audio and torch are optional voice-fallback dependencies and are not in requirements.txt.",
        "Document parsing — python-docx, pdf2docx, and PyMuPDF for PDF/DOCX/image extraction; Google Vision for OCR on scanned filings where required.",
        "Template rendering — docxtpl (Jinja2-placeholder DOCX renderer) for the drafting assistant's byte-exact template fill (see §5.5).",
        "Transport — Socket.IO via the python-socketio async client, over HTTPS; the AI service is a client only, never a server.",
        "Logging and observability — Loguru with a Google Cloud Logging shipper; the Backend's request ID is propagated end-to-end so any AI action can be correlated back to the originating request.",
    ]
    for bullet in tech_stack_bullets:
        last = _insert_paragraph_after(last, bullet, style_source=bullet_template)

    # ===== §4.3.2 Agents and Orchestration ==================================
    last = _insert_paragraph_after(
        last, "4.3.2 Agents and Orchestration", style_source=heading2_template
    )
    last = _insert_paragraph_after(
        last,
        "The AI service is a multi-agent pipeline. Each agent has a single responsibility, a defined input/output schema, and is gated by the role lock described first below. No agent talks to the lawyer's browser directly — all output flows back through the Backend.",
        style_source=normal_template,
    )
    agent_bullets = [
        "Role Agent (LangGraph) — non-bypassable gate at workspace init; validates and locks the workspace role (ADJUDICATOR, PETITIONER, RESPONDENT, MEDIATOR). Subsequent role changes invalidate all role-derived intelligence so it is re-generated cleanly rather than mixed across roles.",
        "Intelligence Chat Agent (LangGraph, six fixed nodes) — context_builder → intent_classifier → explain_why (Gemini) → action_dispatcher → response_formatter → governance_validator. Replies are deliberately role-neutral and accuracy-first; state-mutating intents (DRAFT, VALIDATE, STRATEGY_CHANGE, REMOVE_INSIGHT) are routed to specialised agents instead of being answered inline.",
        "Document Assessor — per-document readability scoring, quality-issue flagging, and evidentiary-strength assessment. Emits the document_assessments feature.",
        "Workspace Metadata Generator — derives case title, jurisdiction, category, and short description from the document corpus. Emits the workspace_metadata feature.",
        "Strategy / Objective Generator — one-shot Gemini call over the concatenated vector-store content; emits three to five role-conditioned strategies as strategy_objective.",
        "Analysis Report Generator — produces the quality score, strengths, weaknesses, and missing-information leads. Dispatched as a background task so chat handling is never blocked.",
        "List-of-Dates Generator — extracts a chronological event timeline from the documents, with each entry carrying the source-document citation.",
        "Applicable-Laws Pipeline — retrieves citations from Indian Kanoon and Vertex AI Search, composes the final list via Gemini, then post-LLM filters and caps the output (see §4.3.5).",
        "Procedural Timeline Generator — infers upcoming filing, hearing, and compliance dates from the record.",
        "Drafting Agent — one-time template annotation followed by per-draft variable fill, rendered through docxtpl into the lawyer's exact DOCX template (see §5.5).",
        "Audio Transcription Agent — Chirp 3 BatchRecognize primary path; Gemini content re-diarization step for two-speaker legal calls to re-attribute utterances by conversational role.",
    ]
    for bullet in agent_bullets:
        last = _insert_paragraph_after(last, bullet, style_source=bullet_template)

    # ===== §4.3.3 Feature Pipeline and Outputs ==============================
    last = _insert_paragraph_after(
        last, "4.3.3 Feature Pipeline and Outputs", style_source=heading2_template
    )
    last = _insert_paragraph_after(
        last,
        "The AI service produces a fixed set of eight feature outputs per workspace. Each output is emitted as a document_feature_response Socket.IO event with a featureType discriminator so the Backend can apply the right merge rules:",
        style_source=normal_template,
    )
    pipeline_bullets = [
        "Trigger events — workspace_created starts the full pipeline; workspace_sync re-runs the affected features when documents are added, deleted, or updated.",
        "featureType values — document_assessments, workspace_metadata, strategy_objective, analysis_report, list_of_dates, applicable_laws, procedural_timeline, drafting_assistant.",
        "Dispatch model — short-running features execute sequentially inside the inbound socket handler; long-running features (analysis_report, applicable_laws, procedural_timeline) are dispatched via asyncio.create_task so chat handling remains responsive.",
        "Replay cache — every emitted payload is held in an in-memory cache keyed by workspace_id (one cache kind per featureType). If the Backend disconnects mid-emit, the AI service re-emits the cached payloads on reconnect so the client never misses a feature update.",
        "Persistence beyond memory — feature snapshots are also persisted into Qdrant as zero-vector payload-only points, so a full service restart still recovers per-workspace feature state without re-running the pipeline.",
    ]
    for bullet in pipeline_bullets:
        last = _insert_paragraph_after(last, bullet, style_source=bullet_template)

    # ===== §4.3.4 Socket Events — Inbound and Outbound ======================
    last = _insert_paragraph_after(
        last,
        "4.3.4 Socket Events — Inbound and Outbound",
        style_source=heading2_template,
    )
    last = _insert_paragraph_after(
        last,
        "The AI service listens for a fixed set of inbound events from the Backend and emits a corresponding set of outbound events. Every event is JSON-validated against the schemas in app/socket/schemas.py before being acted on or delivered.",
        style_source=normal_template,
    )
    last = _insert_paragraph_after(
        last,
        "Inbound (Backend → AI service):",
        style_source=normal_template,
    )
    inbound_bullets = [
        "workspace_created — kicks off the full feature pipeline for a newly-created workspace.",
        "workspace_sync — re-runs the affected features when the workspace's document set changes.",
        "user_message_request — a chat turn; runs the intelligence-chat LangGraph.",
        "probing_prompts_request — generates role-aware follow-up prompt suggestions.",
        "analysis_report_feedback and added_missing_information — recalculate the quality score and suppress lawyer-rejected items.",
        "applicable_law_relevance_request — assesses the relevance of a specific law to the current case facts.",
        "analysis_docx_generate_request — normalises the analysis report for DOCX export.",
        "template_sync — caches or evicts the lawyer's drafting template library so the drafting agent stays in sync with the lawyer's uploads.",
        "ai_heartbeat_ping — periodic liveness probe from the Backend.",
    ]
    for bullet in inbound_bullets:
        last = _insert_paragraph_after(last, bullet, style_source=bullet_template)

    last = _insert_paragraph_after(
        last,
        "Outbound (AI service → Backend):",
        style_source=normal_template,
    )
    outbound_bullets = [
        "document_feature_response — carries one of the eight featureType payloads described in §4.3.3.",
        "ai_response — the chat reply emitted by the intelligence-chat agent.",
        "probing_prompts_response — the paired reply to probing_prompts_request.",
        "applicable_law_relevance_response — the paired reply to applicable_law_relevance_request.",
        "analysis_docx_generate_response — the normalised analysis-report payload ready for DOCX export.",
        "ai_heartbeat_pong — liveness reply to the Backend's heartbeat ping.",
    ]
    for bullet in outbound_bullets:
        last = _insert_paragraph_after(last, bullet, style_source=bullet_template)

    # ===== Rename §4.3.1 Applicable-Laws … → §4.3.5 Indian-Law Grounding ====
    _replace_paragraph_text(
        existing_4_3_1,
        "4.3.5 Indian-Law Grounding and Applicable-Laws Sourcing",
    )

    doc.save(str(DOCX_PATH))
    print(f"Expanded §4.3 in {DOCX_PATH.name}")
    print(
        f"  inserted §4.3.1 Technology Stack ({len(tech_stack_bullets)} bullets)\n"
        f"  inserted §4.3.2 Agents and Orchestration ({len(agent_bullets)} bullets)\n"
        f"  inserted §4.3.3 Feature Pipeline and Outputs ({len(pipeline_bullets)} bullets)\n"
        f"  inserted §4.3.4 Socket Events ({len(inbound_bullets)} inbound, {len(outbound_bullets)} outbound)\n"
        f"  renamed §4.3.1 Applicable-Laws → §4.3.5 Indian-Law Grounding"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
