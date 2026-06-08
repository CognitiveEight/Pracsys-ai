#!/usr/bin/env python3
"""Fix the Logging & Observability mix-up in Pracsys_workflow_overview.docx.

`businessify_pracsys_ai_content.py` searched for "Logging & Observability:"
to rewrite the §4.3.1 AI-stack bullet, but the same label exists in §4.2
(the Backend stack) and appears first, so the wrong paragraph was
rewritten. This restores §4.2 to its original wording and applies the
intended business rewrite to §4.3.1.

Run from the repo root:

    python3 scripts/fix_logging_bullet.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(__file__).resolve().parent.parent / "Pracsys_workflow_overview.docx"


def _normalize(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().lower()


def _find_after(doc, prefix: str, *, after: Paragraph):
    """Find the first paragraph after `after` whose text starts with `prefix`."""
    needle = _normalize(prefix)
    seen_anchor = False
    for p in doc.paragraphs:
        if p._p is after._p:
            seen_anchor = True
            continue
        if not seen_anchor:
            continue
        if _normalize(p.text).startswith(needle):
            return p
    raise LookupError(f"not found after anchor: {prefix!r}")


def _find_first(doc, prefix: str):
    needle = _normalize(prefix)
    for p in doc.paragraphs:
        if _normalize(p.text).startswith(needle):
            return p
    raise LookupError(f"not found: {prefix!r}")


def _set_run_text(run_elem, text: str) -> None:
    for t in list(run_elem.findall(qn("w:t"))):
        run_elem.remove(t)
    new_t = OxmlElement("w:t")
    new_t.text = text
    new_t.set(qn("xml:space"), "preserve")
    run_elem.append(new_t)


def replace_two_run_bullet(p: Paragraph, label: str, rest: str) -> None:
    runs = p._p.findall(qn("w:r"))
    if len(runs) >= 2:
        _set_run_text(runs[0], f"{label}: ")
        _set_run_text(runs[1], rest)
        for r in runs[2:]:
            for t in list(r.findall(qn("w:t"))):
                r.remove(t)
    else:
        if not runs:
            p.add_run(f"{label}: {rest}")
        else:
            _set_run_text(runs[0], f"{label}: {rest}")


def main() -> int:
    doc = Document(str(DOCX_PATH))

    # 1) Restore §4.2 Backend Logging & Observability bullet.
    #    Right now that paragraph reads "Observability: Centralised log shipping..."
    #    and is the FIRST "Observability:"-prefixed paragraph in the doc.
    backend_heading = _find_first(doc, "4.2 Backend")
    corrupted = _find_after(doc, "Observability:", after=backend_heading)
    replace_two_run_bullet(
        corrupted,
        "Logging & Observability",
        "Winston with Google Cloud Logging integration; request-ID correlation across every layer",
    )
    print("restored §4.2 Logging & Observability bullet")

    # 2) Rewrite the §4.3.1 bullet that still says "Logging & Observability: Loguru…"
    ai_section_heading = _find_first(doc, "4.3.1 Technology Stack")
    leaky = _find_after(doc, "Logging & Observability:", after=ai_section_heading)
    replace_two_run_bullet(
        leaky,
        "Observability",
        "Centralised log shipping with request-ID correlation across services.",
    )
    print("rewrote §4.3.1 Observability bullet (Loguru reference dropped)")

    doc.save(str(DOCX_PATH))
    return 0


if __name__ == "__main__":
    sys.exit(main())
