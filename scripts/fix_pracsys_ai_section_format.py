#!/usr/bin/env python3
"""Fix formatting of §4.3.1–§4.3.4 in Pracsys_workflow_overview.docx.

The previous expansion (`expand_pracsys_ai_section.py`) ended up with every
bullet rendered fully bold because the helper copied the run-properties of
the source bullet's first run, which happened to be bold ("Boundary:",
"AI Stack:", etc.) and then carried that bold flag onto the single run it
created for the new paragraph.

This script replaces the buggy content of §4.3.1–§4.3.4 with cleanly
formatted bullets that match the §4.1 / §4.2 visual style:

    • <bold label>: <regular description>

It does this by cloning paragraphs that already have the desired two-run
structure (one bold run for the label, one regular run for the rest) and
rewriting the text of each run in place. Run-level formatting comes from
the clone, so no manual bold / font / size handling is needed.

Idempotent for the user request as long as §4.3.5 exists; the script
detects that pre-condition and refuses to run otherwise. If re-run after
its own success, it simply re-applies the same rewrites — they are
content-equivalent so the file will round-trip stably.

Run from the repo root:

    python3 scripts/fix_pracsys_ai_section_format.py
"""

from __future__ import annotations

import re
import sys
from copy import deepcopy
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(__file__).resolve().parent.parent / "Pracsys_workflow_overview.docx"


def _normalize(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().lower()


def _find_paragraph_starting_with(doc, prefix: str) -> Paragraph:
    needle = _normalize(prefix)
    for para in doc.paragraphs:
        if _normalize(para.text).startswith(needle):
            return para
    raise LookupError(f"Could not locate paragraph starting with: {prefix!r}")


def _find_paragraph_text_equals(doc, text: str) -> Paragraph:
    needle = _normalize(text)
    for para in doc.paragraphs:
        if _normalize(para.text) == needle:
            return para
    raise LookupError(f"Could not locate paragraph with exact text: {text!r}")


def _set_run_text(run_elem, text: str) -> None:
    """Replace all w:t children in a run with a single one carrying `text`."""
    for t in list(run_elem.findall(qn("w:t"))):
        run_elem.remove(t)
    new_t = OxmlElement("w:t")
    new_t.text = text
    new_t.set(qn("xml:space"), "preserve")
    run_elem.append(new_t)


def _clone_paragraph_after(
    anchor: Paragraph,
    template: Paragraph,
    *,
    texts: Sequence[str],
) -> Paragraph:
    """Clone `template`'s paragraph element after `anchor` and overwrite text.

    `texts` is a list of per-run replacement strings. If the template has
    more runs than `texts`, extra runs are cleared. If `texts` has more
    entries than the template's runs, the trailing strings are concatenated
    into the last run (so multi-run templates always survive).
    """
    new_p = deepcopy(template._p)
    runs = new_p.findall(qn("w:r"))
    if not runs:
        # Should never happen for the templates we pick, but degrade gracefully.
        r = OxmlElement("w:r")
        _set_run_text(r, "".join(texts))
        new_p.append(r)
    else:
        for i, run_elem in enumerate(runs):
            if i < len(texts):
                _set_run_text(run_elem, texts[i])
            else:
                # Clear any extra runs to avoid leaking template content.
                for t in list(run_elem.findall(qn("w:t"))):
                    run_elem.remove(t)
        if len(texts) > len(runs):
            # Append overflow to the last existing run.
            tail = "".join(texts[len(runs):])
            existing_t = runs[-1].find(qn("w:t"))
            if existing_t is not None:
                existing_t.text = (existing_t.text or "") + tail
                existing_t.set(qn("xml:space"), "preserve")
            else:
                _set_run_text(runs[-1], tail)
    anchor._p.addnext(new_p)
    return Paragraph(new_p, anchor._parent)


def _delete_between(start_exclusive: Paragraph, end_exclusive: Paragraph) -> int:
    """Remove every paragraph strictly between two anchors. Returns count."""
    body = start_exclusive._p.getparent()
    cur = start_exclusive._p.getnext()
    removed = 0
    while cur is not None and cur is not end_exclusive._p:
        nxt = cur.getnext()
        body.remove(cur)
        cur = nxt
        removed += 1
    return removed


def main() -> int:
    if not DOCX_PATH.exists():
        print(f"ERROR: missing {DOCX_PATH}", file=sys.stderr)
        return 1

    doc = Document(str(DOCX_PATH))

    # Locate the bracketing anchors.
    boundary = _find_paragraph_starting_with(doc, "Boundary: The AI service never")
    section_4_3_5 = _find_paragraph_starting_with(doc, "4.3.5 Indian-Law Grounding")

    # Style templates we'll clone.
    framework_bullet = _find_paragraph_starting_with(doc, "Framework: Next.js")  # bold-prefix bullet
    heading2_template = section_4_3_5                                            # Heading 2
    sensitive_fields_header = _find_paragraph_text_equals(doc, "Sensitive Fields:")  # bold Normal

    # Wipe out everything between Boundary (the last §4.3 high-level bullet)
    # and §4.3.5 (the renamed Applicable-Laws heading). That removes the
    # buggy §4.3.1–§4.3.4 we inserted previously.
    removed = _delete_between(boundary, section_4_3_5)
    print(f"removed {removed} previously-inserted paragraphs in §4.3.1–§4.3.4")

    # ---- Helpers bound to our chosen templates ----------------------------
    def insert_heading(anchor: Paragraph, text: str) -> Paragraph:
        # Heading 2 paragraphs typically have a single run; one text slot.
        return _clone_paragraph_after(anchor, heading2_template, texts=[text])

    def insert_bullet(anchor: Paragraph, label: str, rest: str) -> Paragraph:
        # Framework template: run[0] BOLD "Framework: ", run[1] regular "Next.js…"
        # We force the same shape: bold "Label: " + regular rest.
        return _clone_paragraph_after(
            anchor,
            framework_bullet,
            texts=[f"{label}: ", rest],
        )

    def insert_subheader(anchor: Paragraph, text: str) -> Paragraph:
        # Sensitive Fields template: one bold run on a Normal paragraph.
        return _clone_paragraph_after(anchor, sensitive_fields_header, texts=[text])

    last = boundary

    # ===== §4.3.1 Technology Stack =========================================
    last = insert_heading(last, "4.3.1 Technology Stack")
    for label, rest in [
        ("Language & Runtime", "Python 3.11+ asyncio service; single long-lived process per environment."),
        ("LLM", "Google Gemini family — used for reasoning, classification, drafting, and audio re-diarization."),
        ("Embeddings", "Google text-embedding-004 — 768-dimensional vectors with cosine similarity."),
        ("Vector Store", "Qdrant — a single collection holds every workspace's document chunks and feature-state snapshots; a workspace_id payload key filters every read and write so workspaces are strictly isolated."),
        ("Agent Orchestration", "LangGraph — state-machine framework used for the role agent and the intelligence-chat graph."),
        ("Law-Citation Retrieval", "Google Vertex AI Search (primary, via langchain_google_community.VertexAISearchRetriever) and Indian Kanoon (indiankanoon.org, site-restricted retrieval) as fallback."),
        ("Speech-to-Text", "Google Speech-to-Text V2 (Chirp 3) with GCS-backed BatchRecognize for native speaker diarization."),
        ("Audio Pre-processing", "ffmpeg for 16-bit PCM / 16 kHz mono normalisation; pyannote.audio and torch are optional voice-fallback dependencies and are not in requirements.txt."),
        ("Document Parsing", "python-docx, pdf2docx, and PyMuPDF for PDF / DOCX / image extraction; Google Vision for OCR on scanned filings where required."),
        ("Template Rendering", "docxtpl (Jinja2-placeholder DOCX renderer) for the drafting assistant's byte-exact template fill (see §5.5)."),
        ("Transport", "Socket.IO via the python-socketio async client over HTTPS; the AI service is a client only, never a server."),
        ("Logging & Observability", "Loguru with a Google Cloud Logging shipper; the Backend's request ID is propagated end-to-end so any AI action can be correlated back to the originating request."),
    ]:
        last = insert_bullet(last, label, rest)

    # ===== §4.3.2 Agents and Orchestration =================================
    last = insert_heading(last, "4.3.2 Agents and Orchestration")
    for label, rest in [
        ("Role Agent", "Non-bypassable LangGraph gate at workspace init; validates and locks the workspace role (ADJUDICATOR, PETITIONER, RESPONDENT, MEDIATOR). Subsequent role changes invalidate all role-derived intelligence so it is re-generated cleanly."),
        ("Intelligence Chat Agent", "LangGraph pipeline with six fixed nodes — context_builder, intent_classifier, explain_why (Gemini), action_dispatcher, response_formatter, governance_validator. Replies are role-neutral and accuracy-first; state-mutating intents (DRAFT, VALIDATE, STRATEGY_CHANGE, REMOVE_INSIGHT) are routed to specialised agents instead of being answered inline."),
        ("Document Assessor", "Per-document readability scoring, quality-issue flagging, and evidentiary-strength assessment. Emits the document_assessments feature."),
        ("Workspace Metadata Generator", "Derives case title, jurisdiction, category, and short description from the document corpus. Emits the workspace_metadata feature."),
        ("Strategy & Objective Generator", "One-shot Gemini call over the concatenated vector-store content; emits three to five role-conditioned strategies as strategy_objective."),
        ("Analysis Report Generator", "Produces the quality score, strengths, weaknesses, and missing-information leads. Dispatched as a background task so chat handling is never blocked."),
        ("List-of-Dates Generator", "Extracts a chronological event timeline with source-document citations. Emits the list_of_dates feature."),
        ("Applicable-Laws Pipeline", "Retrieves citations from Indian Kanoon and Vertex AI Search, composes the final list via Gemini, then post-LLM filters and caps the output (see §4.3.5)."),
        ("Procedural Timeline Generator", "Infers upcoming filing, hearing, and compliance dates from the record. Emits the procedural_timeline feature."),
        ("Drafting Agent", "One-time template annotation followed by per-draft variable fill, rendered through docxtpl into the lawyer's exact DOCX template (see §5.5). Emits the drafting_assistant feature."),
        ("Audio Transcription Agent", "Chirp 3 BatchRecognize primary path; Gemini content re-diarization for two-speaker legal calls to re-attribute utterances by conversational role."),
    ]:
        last = insert_bullet(last, label, rest)

    # ===== §4.3.3 Feature Pipeline and Outputs =============================
    last = insert_heading(last, "4.3.3 Feature Pipeline and Outputs")
    for label, rest in [
        ("Trigger Events", "workspace_created starts the full pipeline; workspace_sync re-runs the affected features when documents are added, deleted, or updated."),
        ("featureType Values", "document_assessments, workspace_metadata, strategy_objective, analysis_report, list_of_dates, applicable_laws, procedural_timeline, drafting_assistant — eight discriminators carried by the document_feature_response event."),
        ("Dispatch Model", "Short-running features execute sequentially inside the inbound socket handler; long-running features (analysis report, applicable laws, procedural timeline) are dispatched via asyncio.create_task so chat handling remains responsive."),
        ("Replay Cache", "Every emitted payload is held in an in-memory cache keyed by workspace_id (one cache kind per featureType). If the Backend disconnects mid-emit, the AI service re-emits the cached payloads on reconnect so the client never misses a feature update."),
        ("Persistent Snapshots", "Feature snapshots are also persisted into Qdrant as zero-vector payload-only points, so a full service restart still recovers per-workspace feature state without re-running the pipeline."),
    ]:
        last = insert_bullet(last, label, rest)

    # ===== §4.3.4 Socket Events — Inbound and Outbound =====================
    last = insert_heading(last, "4.3.4 Socket Events — Inbound and Outbound")
    last = insert_subheader(last, "Inbound (Backend → AI service):")
    for label, rest in [
        ("workspace_created", "Kicks off the full feature pipeline for a newly-created workspace."),
        ("workspace_sync", "Re-runs the affected features when the workspace's document set changes."),
        ("user_message_request", "A chat turn; runs the intelligence-chat LangGraph."),
        ("probing_prompts_request", "Generates role-aware follow-up prompt suggestions."),
        ("analysis_report_feedback / added_missing_information", "Recalculate the quality score and suppress lawyer-rejected items."),
        ("applicable_law_relevance_request", "Assesses the relevance of a specific law to the current case facts."),
        ("analysis_docx_generate_request", "Normalises the analysis report for DOCX export."),
        ("template_sync", "Caches or evicts the lawyer's drafting template library so the drafting agent stays in sync with the lawyer's uploads."),
        ("ai_heartbeat_ping", "Periodic liveness probe from the Backend."),
    ]:
        last = insert_bullet(last, label, rest)
    last = insert_subheader(last, "Outbound (AI service → Backend):")
    for label, rest in [
        ("document_feature_response", "Carries one of the eight featureType payloads described in §4.3.3."),
        ("ai_response", "The chat reply emitted by the intelligence-chat agent."),
        ("probing_prompts_response", "Paired reply to probing_prompts_request."),
        ("applicable_law_relevance_response", "Paired reply to applicable_law_relevance_request."),
        ("analysis_docx_generate_response", "Normalised analysis-report payload ready for DOCX export."),
        ("ai_heartbeat_pong", "Liveness reply to the Backend's heartbeat ping."),
    ]:
        last = insert_bullet(last, label, rest)

    doc.save(str(DOCX_PATH))
    print(f"Rewrote §4.3.1–§4.3.4 in {DOCX_PATH.name} with §4.1/§4.2-style formatting")
    return 0


if __name__ == "__main__":
    sys.exit(main())
