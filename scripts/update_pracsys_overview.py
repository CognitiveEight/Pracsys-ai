#!/usr/bin/env python3
"""One-shot rewrite of AI-related sections in Pracsys_workflow_overview.docx.

Rewrites only the AI-related paragraphs (§1 feature list, §4.3 AI service,
§5.3 Step 3 & 4, §5.4 intelligence chat, §6.2 audio files, §7.2 backend↔AI
auth) to match the actual SimplerToday Python AI service implementation,
and adds two new subsections (§4.3.1 applicable-laws sourcing, §5.5
drafting assistant) plus five glossary entries. Non-AI content is left
untouched.

Run from the repo root:

    python3 scripts/update_pracsys_overview.py
"""

from __future__ import annotations

import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(__file__).resolve().parent.parent / "Pracsys_workflow_overview.docx"


def _normalize(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().lower()


def _find_paragraph_starting_with(doc, prefix: str) -> Paragraph:
    needle = _normalize(prefix)
    for para in doc.paragraphs:
        if _normalize(para.text).startswith(needle):
            return para
    raise LookupError(f"Could not locate paragraph starting with: {prefix!r}")


def _replace_paragraph_text(paragraph: Paragraph, new_text: str) -> None:
    """Rewrite the entire paragraph text, keeping the first run's formatting.

    Style (List Bullet, Normal, Heading 2, etc.) is unaffected because it
    lives on the paragraph's w:pPr, not on the runs.
    """
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def _replace_paragraph_substring(
    paragraph: Paragraph, old_substring: str, new_substring: str
) -> bool:
    """Substring replacement that survives multi-run paragraphs.

    Returns True if the substring was found and replaced.
    """
    full = "".join(r.text for r in paragraph.runs)
    if old_substring not in full:
        return False

    start = full.find(old_substring)
    end = start + len(old_substring)
    spans = []
    cursor = 0
    for run in paragraph.runs:
        spans.append((cursor, cursor + len(run.text), run))
        cursor += len(run.text)
    affected = [(s, e, r) for (s, e, r) in spans if not (e <= start or s >= end)]
    if not affected:
        return False

    if len(affected) == 1:
        s, _e, r = affected[0]
        local_s = start - s
        local_e = end - s
        r.text = r.text[:local_s] + new_substring + r.text[local_e:]
        return True

    first_s, _, first_r = affected[0]
    last_s, _, last_r = affected[-1]
    first_local = start - first_s
    last_local = end - last_s
    first_r.text = first_r.text[:first_local] + new_substring
    for (_, _, r) in affected[1:-1]:
        r.text = ""
    last_r.text = last_r.text[last_local:]
    return True


def _insert_paragraph_after(
    anchor: Paragraph, text: str, *, style_source: Paragraph
) -> Paragraph:
    """Insert a new paragraph with `text` immediately after `anchor`.

    Style and run-properties are inherited from `style_source` by
    deep-copying its w:p element (which carries the w:pPr style ref) and
    its first run's w:rPr.
    """
    new_p = deepcopy(style_source._p)
    # Strip existing runs from the copy.
    for r in list(new_p.findall(qn("w:r"))):
        new_p.remove(r)
    # Capture run formatting (font, size, bold, italic) from the source.
    rpr_copy = None
    src_first_r = style_source._p.find(qn("w:r"))
    if src_first_r is not None:
        rpr = src_first_r.find(qn("w:rPr"))
        if rpr is not None:
            rpr_copy = deepcopy(rpr)
    # Build a fresh run with the new text.
    new_r = OxmlElement("w:r")
    if rpr_copy is not None:
        new_r.append(rpr_copy)
    new_t = OxmlElement("w:t")
    new_t.text = text
    new_t.set(qn("xml:space"), "preserve")
    new_r.append(new_t)
    new_p.append(new_r)

    anchor._p.addnext(new_p)
    return Paragraph(new_p, anchor._parent)


def main() -> int:
    if not DOCX_PATH.exists():
        print(f"ERROR: missing {DOCX_PATH}", file=sys.stderr)
        return 1

    doc = Document(str(DOCX_PATH))

    # ----- Locate style-source paragraphs (used as templates for inserts) ---
    heading2_template = _find_paragraph_starting_with(doc, "4.3 AI Processing Service")
    normal_template = _find_paragraph_starting_with(doc, "Each workspace becomes")

    # ===== §1 Summary — feature bullets =====================================
    s1_analysis = _find_paragraph_starting_with(doc, "Analysis Reports")
    s1_dates = _find_paragraph_starting_with(doc, "List of Dates")
    s1_laws = _find_paragraph_starting_with(doc, "Applicable Laws")
    s1_parties = _find_paragraph_starting_with(doc, "Relevant Parties")
    s1_drafting = _find_paragraph_starting_with(doc, "Drafting Assistant")
    s1_support = _find_paragraph_starting_with(doc, "Supporting Material")

    _replace_paragraph_text(
        s1_analysis,
        "Document Assessments — per-document readability and evidentiary-strength scoring, with quality issues flagged so weaker evidence is visible before it reaches a filing.",
    )
    _replace_paragraph_text(
        s1_dates,
        "Workspace Metadata — derived case title, jurisdiction, category, and short description used to organise the lawyer's case list.",
    )
    _replace_paragraph_text(
        s1_laws,
        "Strategy & Objective — three to five role-conditioned strategies or opinions per workspace, refined as more evidence is uploaded.",
    )
    _replace_paragraph_text(
        s1_parties,
        "Analysis Report — a professional, branded case analysis (quality score, strengths, weaknesses, missing-information leads) suitable for client and counsel review.",
    )
    _replace_paragraph_text(
        s1_drafting,
        "List of Dates — a chronologically organised, source-cited event timeline.",
    )
    _replace_paragraph_text(
        s1_support,
        "Applicable Laws — Indian-law citations (BNS, BNSS, BSA, IT Act and the wider statutory set) and grounded case law, mapped directly to the uploaded evidence.",
    )
    after_support_1 = _insert_paragraph_after(
        s1_support,
        "Procedural Timeline — upcoming filing, hearing, and compliance dates inferred from the record so deadlines are not missed.",
        style_source=s1_support,
    )
    _insert_paragraph_after(
        after_support_1,
        "Drafting Assistant — AI-suggested document drafts (notices, applications, replies) rendered into the lawyer's own DOCX template, with logo, fonts, headers and statutory boilerplate preserved byte-for-byte.",
        style_source=s1_support,
    )

    # ===== §4.3 AI Processing Service ========================================
    s43_runtime = _find_paragraph_starting_with(doc, "Runtime: Independent Python service")
    s43_channel = _find_paragraph_starting_with(doc, "Channel: Socket.IO")
    s43_resp = _find_paragraph_starting_with(doc, "Responsibilities:")
    s43_boundary = _find_paragraph_starting_with(doc, "Boundary:")

    _replace_paragraph_text(
        s43_runtime,
        "Runtime: Independent Python (asyncio) service deployed in isolation from the API layer; at boot it initialises the Qdrant vector store, the document processor, the Google Gemini client, and the Socket.IO client in that order, then blocks on a shutdown signal.",
    )
    _replace_paragraph_text(
        s43_channel,
        "Channel: Socket.IO over an authenticated, server-to-server connection (HMAC-SHA256 signature of a service claim signed with a shared secret, sent in the connection headers; see §7.2).",
    )
    _replace_paragraph_text(
        s43_resp,
        "Responsibilities: Document ingestion and OCR, audio transcription with speaker diarization, classification, entity extraction, chronology construction, applicable-law and case-law retrieval, draft generation into lawyer-supplied DOCX templates.",
    )
    # Insert AI Stack bullet AFTER channel, BEFORE responsibilities.
    _insert_paragraph_after(
        s43_channel,
        "AI Stack: Google Gemini (LLM for reasoning and drafting; text-embedding-004 for 768-dimension cosine-similarity embeddings) for all language work; Qdrant as the vector store, scoped per workspace via a workspace_id payload filter on every read and write; LangGraph as the agent-orchestration framework (role agent and intelligence-chat graph); Google Vertex AI Search and Indian Kanoon (indiankanoon.org) as grounded sources for applicable-laws output; Google Speech-to-Text V2 (Chirp 3) for audio evidence transcription.",
        style_source=s43_channel,
    )

    # ===== New §4.3.1 — Applicable-Laws Sourcing =============================
    # Insert immediately AFTER the Boundary bullet so it appears at the end
    # of §4.3 but before §5 starts.
    last = s43_boundary
    last = _insert_paragraph_after(
        last,
        "4.3.1 Applicable-Laws Sourcing and Indian-Law Guardrails",
        style_source=heading2_template,
    )
    last = _insert_paragraph_after(
        last,
        "Applicable-laws output is grounded against live Indian-law sources rather than relying on the LLM's pre-training alone. Two retrievers feed the AI before Gemini composes the final citation list:",
        style_source=normal_template,
    )
    last = _insert_paragraph_after(
        last,
        "Indian Kanoon (indiankanoon.org) — the primary direct lookup, run against the largest open Indian case-law and statutes corpus; queried with a site-restricted search.",
        style_source=s43_boundary,
    )
    last = _insert_paragraph_after(
        last,
        "Google Vertex AI Search — the fallback / supplementary retriever, queried via langchain_google_community.VertexAISearchRetriever against a curated data store (GCP_PROJECT_ID and VERTEX_SEARCH_DATA_STORE_ID). Vertex AI Search replaces the previously-used Google Custom Search integration.",
        style_source=s43_boundary,
    )
    last = _insert_paragraph_after(
        last,
        "Two hard constraints govern the final emitted list:",
        style_source=normal_template,
    )
    last = _insert_paragraph_after(
        last,
        "Indian Penal Code, Code of Criminal Procedure, and Indian Evidence Act are filtered out of the AI output by a post-LLM normaliser. The Bharatiya Nyaya Sanhita (BNS), Bharatiya Nagarik Suraksha Sanhita (BNSS), and Bharatiya Sakshya Adhiniyam (BSA) are the operative statutes for post-1 July 2024 matters and are the only criminal-law framing allowed in output.",
        style_source=s43_boundary,
    )
    last = _insert_paragraph_after(
        last,
        "The applicable-laws list is capped at twelve items with stable identifiers law_1 through law_12 assigned in discovery order, so the lawyer's highlights and notes on a given law slot survive re-runs.",
        style_source=s43_boundary,
    )

    # ===== §5.3 Step 3 — AI bootstrapping ====================================
    step3 = _find_paragraph_starting_with(doc, "Step 3 — AI bootstrapping")
    _replace_paragraph_text(
        step3,
        "Step 3 — AI bootstrapping. The Backend emits a workspace_created event over the server-to-server Socket.IO channel to the AI service. The payload carries the workspace ID, owner ID, title, the workspace role (one of ADJUDICATOR, PETITIONER, RESPONDENT, MEDIATOR), and short-lived signed URLs for the evidence files. The AI service routes the request through a non-bypassable role gate implemented in LangGraph: the role is validated and locked before any downstream feature agent — analysis, strategy, list of dates, applicable laws, procedural timeline, drafting — is allowed to run. Subsequent role changes invalidate all derived intelligence so it is re-generated cleanly rather than mixed across roles. The user does not need to wait — processing is asynchronous and the lawyer can begin interacting with the workspace immediately.",
    )

    # ===== §5.3 Step 4 — Progressive enrichment ==============================
    step4 = _find_paragraph_starting_with(doc, "Step 4 — Progressive enrichment")
    old_step4_tail = (
        " As the AI completes each analytical pass, it emits feature-specific "
        "responses back to the Backend — analysis report, list of dates, "
        "applicable laws, relevant parties, case laws, procedural timeline, "
        "drafting assistant, strategy and objective. The Backend:"
    )
    new_step4_tail = (
        " As the AI completes each analytical pass it emits a "
        "document_feature_response event back to the Backend tagged with one "
        "of eight featureType values: document_assessments, "
        "workspace_metadata, strategy_objective, analysis_report, "
        "list_of_dates, applicable_laws, procedural_timeline, and "
        "drafting_assistant. Long-running passes (analysis report, applicable "
        "laws, procedural timeline) are dispatched via asyncio.create_task so "
        "chat handling is never blocked, and every emitted payload is held in "
        "an in-memory replay cache keyed by workspace_id so the AI service "
        "can re-emit it after a socket drop without re-running the full "
        "pipeline. The Backend:"
    )
    if not _replace_paragraph_substring(step4, old_step4_tail, new_step4_tail):
        _replace_paragraph_text(step4, "Step 4 — Progressive enrichment." + new_step4_tail)

    # ===== §5.4 Intelligence Chat ============================================
    chat_intro = _find_paragraph_starting_with(doc, "Inside any workspace the lawyer can ask")
    _replace_paragraph_text(
        chat_intro,
        "Inside any workspace the lawyer can ask natural-language questions "
        "(\"Summarise the prosecution's argument\", \"Draft a reply to the "
        "notice dated 12 March\"). The AI side is implemented as a LangGraph "
        "pipeline with six fixed nodes — context_builder → intent_classifier "
        "→ explain_why (Gemini) → action_dispatcher → response_formatter → "
        "governance_validator — gated by the same role lock as the rest of "
        "the AI service. The flow:",
    )
    chat_forward = _find_paragraph_starting_with(doc, "The Backend persists the message")
    _replace_paragraph_text(
        chat_forward,
        "The Backend persists the message, then forwards it to the AI "
        "service. The AI service assembles a ChatContext containing the "
        "workspace role, document snippets retrieved from Qdrant (capped per "
        "document and overall to stay within the model's context budget), "
        "the active strategy, approved facts, the current analysis-report "
        "snapshot, and any inline attachment, then runs it through the "
        "LangGraph chat graph.",
    )
    chat_audit = _find_paragraph_starting_with(doc, "Every chat turn is durably stored")
    _replace_paragraph_text(
        chat_audit,
        "Every chat turn is durably stored against the workspace and "
        "contributes to the workspace's audit history. Chat replies are "
        "deliberately role-neutral and accuracy-first — role-conditioned "
        "advocacy lives in the panel outputs (strategy, drafts), not in the "
        "chat answer. The action_dispatcher node only labels which "
        "downstream agent should service an intent (for example, DRAFT → "
        "drafting agent, STRATEGY_CHANGE → strategy agent); it never invokes "
        "other agents inline. The governance_validator node blocks "
        "state-mutating intents (DRAFT, VALIDATE, STRATEGY_CHANGE, "
        "REMOVE_INSIGHT) from being answered directly and ensures they are "
        "routed to the appropriate agent instead.",
    )

    # ===== New §5.5 — Drafting Assistant =====================================
    last = chat_audit
    last = _insert_paragraph_after(
        last,
        "5.5 Drafting Assistant — Exact Template Preservation",
        style_source=heading2_template,
    )
    last = _insert_paragraph_after(
        last,
        "The Drafting Assistant lets a lawyer convert any of their own DOCX templates (notices, replies, applications) into a reusable, AI-fillable draft without compromising the firm's letterhead, fonts, logo, headers/footers, or statutory boilerplate.",
        style_source=normal_template,
    )
    last = _insert_paragraph_after(
        last,
        "Step 1 — Template ingestion (one time per template). Gemini reads the template text and emits a strict JSON list of variable regions: recipient name, case number, dates, factual narrative, prayer wording, signatures, amounts. Firm name, advocate roll number, statutory citation text, sign-offs, and other template-fixed content are explicitly excluded by the annotation prompt's rules.",
        style_source=normal_template,
    )
    last = _insert_paragraph_after(
        last,
        "Step 2 — In-place Jinja2 annotation. The annotation routine locates each variable region across body paragraphs, tables, and headers/footers, then mutates the matching run's text to a Jinja2 placeholder ({{ var_name }}) without recreating the run object. Font, size, bold, italic, colour, and spacing survive untouched and the surrounding OOXML structure is preserved.",
        style_source=normal_template,
    )
    last = _insert_paragraph_after(
        last,
        "Step 3 — Per-draft fill. For each new draft, Gemini fills the variable schema against the case facts (Indian-law grounded, named parties, dates in DD/MM/YYYY by default) and the docxtpl renderer injects the values into the annotated template. The output is byte-identical to the lawyer's original template except for the variable text.",
        style_source=normal_template,
    )
    last = _insert_paragraph_after(
        last,
        "Step 4 — Storage and delivery. The rendered DOCX is uploaded to Google Cloud Storage; the Backend returns a short-lived pre-signed URL to the frontend, identical to the evidence-file access pattern described in §6.2.",
        style_source=normal_template,
    )

    # ===== §6.2 Audio Files ==================================================
    audio = _find_paragraph_starting_with(doc, "Audio Files:")
    _replace_paragraph_text(
        audio,
        "Audio Files: Audio uploads (mp3, wav, m4a, ogg, flac, opus, webm, "
        "aac) are transcribed by the AI service using Google Speech-to-Text "
        "V2 (Chirp 3). The primary path is GCS-backed BatchRecognize with "
        "native speaker diarization; two-speaker legal calls additionally "
        "pass through a Gemini content-based re-diarization step that "
        "re-attributes utterances by conversational role. The transcript is "
        "rendered as [MM:SS] Speaker A: text lines and becomes the "
        "document's extracted_text, so quality scoring, vector storage, and "
        "chat context work without special-casing. A transcriptStatus field "
        "(ready / pending / failed) is attached per audio entry. Both the "
        "source audio and the generated transcript are stored, with explicit "
        "linkage via the source file's hash so provenance is preserved. Full "
        "diarization requires ffmpeg on the AI service host and a configured "
        "GCS bucket; without GCS, transcription still works but speaker tags "
        "may be absent.",
    )

    # ===== §7.2 Backend ↔ AI Service authentication ==========================
    bk_auth = _find_paragraph_starting_with(doc, "Authentication: The handshake")
    _replace_paragraph_text(
        bk_auth,
        "Authentication: The Socket.IO handshake carries an HMAC-SHA256 "
        "signature of a server-side claim payload, signed with a shared "
        "secret and base64 URL-safe encoded, sent as the x-signature "
        "header alongside an x-client-type identifier. The AI service is "
        "not exposed to the public internet; only the Backend is "
        "whitelisted to connect.",
    )

    # ===== §11 Glossary — add AI-stack entries ===============================
    last_gloss = _find_paragraph_starting_with(doc, "Socket.IO —")
    last = last_gloss
    last = _insert_paragraph_after(
        last,
        "Google Gemini — Google's family of large language models. SimplerToday's AI service uses Gemini for reasoning, drafting, and (via text-embedding-004) document embeddings.",
        style_source=last_gloss,
    )
    last = _insert_paragraph_after(
        last,
        "Qdrant — Open-source vector database used to store document chunk embeddings. Scoped per workspace via a workspace_id payload filter on every read and write.",
        style_source=last_gloss,
    )
    last = _insert_paragraph_after(
        last,
        "LangGraph — A state-machine framework for orchestrating LLM agents. SimplerToday uses it for the role agent and the intelligence-chat pipeline.",
        style_source=last_gloss,
    )
    last = _insert_paragraph_after(
        last,
        "Vertex AI Search — Google Cloud's managed retrieval service. SimplerToday uses it (alongside Indian Kanoon) as a grounded source for applicable-laws citations.",
        style_source=last_gloss,
    )
    last = _insert_paragraph_after(
        last,
        "Chirp 3 — Google's foundation speech-recognition model, exposed via Speech-to-Text V2. SimplerToday uses it for audio evidence transcription with native speaker diarization.",
        style_source=last_gloss,
    )

    doc.save(str(DOCX_PATH))
    print(f"Updated {DOCX_PATH.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
