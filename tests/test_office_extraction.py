"""Excel (.xlsx) and PowerPoint (.pptx) evidence must be read, not misrouted.

The backend's ``filename`` carries no suffix ("annexure-3-age-verification-
test-results"); ``extension`` is the only type signal. Before this support,
every such file fell into content sniffing, where any ZIP was handed to the
DOCX parser, so spreadsheets and decks came out as an extraction error.
"""

import asyncio
import datetime as dt
from io import BytesIO

import pytest

openpyxl = pytest.importorskip("openpyxl")
pptx = pytest.importorskip("pptx")

from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.processing.file_handlers import (  # noqa: E402
    FileProcessor,
    _zip_office_kind,
    filename_for_extraction,
)
from app.rag.chunker import PAGE_MARKER_RE  # noqa: E402


def _workbook_bytes(rows_on_results: int = 3) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    # Table starts at column B: the empty leading column must be dropped.
    ws.append([None, "Player", "Date of test", "Bone age", "Eligible"])
    for i in range(rows_on_results):
        ws.append([None, f"Player {i + 1}", dt.datetime(2026, 9, 1 + i % 28), 15.5, i % 2 == 0])
    ws.append([])  # blank row: skipped
    ws.append([None, "Total", None, "=SUM(D2:D4)"])  # formula, no cached value
    hidden = wb.create_sheet("Internal notes")
    hidden.sheet_state = "hidden"
    hidden.append(["Remark", "Re-test ordered for Player 2"])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _deck_bytes() -> bytes:
    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    s1.shapes.title.text = "Selection criteria"
    s1.placeholders[1].text = "Age verified by TW3 method\nMinimum 3 domestic matches"
    s1.notes_slide.notes_text_frame.text = "Presented by the convenor"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s2.shapes.title.text = "Trial scores"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Player", "Score"
    table.cell(1, 0).text, table.cell(1, 1).text = "Player 1", "87"
    data = CategoryChartData()
    data.categories = ["Batting", "Bowling"]
    data.add_series("Player 1", (42, 7))
    s2.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(4), Inches(2), data
    )

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Withdrawn proposal"
    s3._element.set("show", "0")  # hidden slide
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def test_extension_is_appended_only_when_missing():
    assert filename_for_extraction("annexure-3-results", "xlsx") == "annexure-3-results.xlsx"
    assert filename_for_extraction("annexure-3-results", ".XLSX") == "annexure-3-results.xlsx"
    assert filename_for_extraction("deck.pptx", "pptx") == "deck.pptx"
    assert filename_for_extraction("deck", None) == "deck"


def test_zip_sniff_tells_office_formats_apart():
    assert _zip_office_kind(_workbook_bytes()) == "xlsx"
    assert _zip_office_kind(_deck_bytes()) == "pptx"
    assert _zip_office_kind(b"PK not really a zip") == "docx"


def test_xlsx_rows_headers_dates_and_hidden_sheets():
    text = FileProcessor.process_file_by_extension(_workbook_bytes(), "results.xlsx", True)
    assert not text.startswith("[")
    assert "Sheet: Results (rows 2-6)" in text
    assert "Player | Date of test | Bone age | Eligible" in text
    assert "Player 1 | 2026-09-01 | 15.5 | TRUE" in text
    assert "Sheet: Internal notes (hidden) (row 1)" in text
    assert "Remark | Re-test ordered for Player 2" in text
    # Leading empty column dropped; a sheet is not a page.
    assert "\n | " not in text and not PAGE_MARKER_RE.search(text)


def test_every_chunk_of_a_long_sheet_carries_its_header():
    from app.rag.chunker import chunk_text

    text = FileProcessor.extract_text_from_xlsx(_workbook_bytes(300))
    chunks = chunk_text(text)
    assert len(chunks) > 3
    header = "Player | Date of test | Bone age | Eligible"
    assert all(header in c.text and "Sheet: Results (rows " in c.text for c in chunks)
    # No row is lost or duplicated across blocks.
    assert all(text.count(f"Player {i} | ") == 1 for i in range(1, 301))


def test_xlsx_cell_cap_is_reported(monkeypatch):
    from app.processing import tables

    monkeypatch.setattr(tables, "MAX_CELLS", 10)
    text = FileProcessor.extract_text_from_xlsx(_workbook_bytes(20))
    assert "Extraction stopped after 10 cells" in text


def test_pptx_slides_tables_charts_notes_and_pages():
    text = FileProcessor.process_file_by_extension(_deck_bytes(), "deck.pptx", True)
    assert not text.startswith("[")
    assert [int(m.group(1)) for m in PAGE_MARKER_RE.finditer(text)] == [1, 2, 3]
    assert "Slide 1: Selection criteria" in text
    assert "Age verified by TW3 method\nMinimum 3 domestic matches" in text
    assert "Speaker notes: Presented by the convenor" in text
    assert "Player 1 | 87" in text
    assert "Player 1: Batting = 42, Bowling = 7" in text
    assert "Slide 3: Withdrawn proposal (hidden)" in text
    # The title is the heading, not repeated as body text.
    assert text.count("Selection criteria") == 1


def test_extensionless_names_are_sniffed_correctly():
    assert "Sheet: Results" in FileProcessor.process_file_by_extension(
        _workbook_bytes(), "annexure-3-age-verification-test-results"
    )
    assert "Slide 1: Selection criteria" in FileProcessor.process_file_by_extension(
        _deck_bytes(), "annexure-4-bcci-selection-committee-presentation"
    )


def test_legacy_formats_fail_loudly_not_as_noise():
    for name in ("old.xls", "old.ppt"):
        out = FileProcessor.process_file_by_extension(b"\xd0\xcf\x11\xe0junk", name)
        assert out.startswith("[Unsupported file format")


def _run_until_routed(monkeypatch, payload: bytes, extension: str):
    """process_document up to the extraction step, recording what it routed to."""
    from app.processing import tables
    from app.processing.document_processor import DocumentProcessor, file_processor

    seen = {}

    async def fake_download(self, url, filename):
        return payload, 1234

    def fake_extract(data, name, mark_pages=False):
        seen["text_path"] = name
        raise RuntimeError("stop after routing")

    def fake_parse(data):
        seen["table_path"] = True
        raise RuntimeError("stop after routing")

    async def fake_classify(self, text, filename):
        # The failure path still classifies the error text; without this stub
        # every test run made real, billed model calls.
        return {"summary": "", "evidentiary_strength": "Low",
                "document_type": "other", "extracted_fields": {}}

    monkeypatch.setattr(DocumentProcessor, "_download_file", fake_download)
    monkeypatch.setattr(DocumentProcessor, "_classify_document", fake_classify)
    monkeypatch.setattr(file_processor, "process_file_by_extension", fake_extract)
    monkeypatch.setattr(tables, "parse_workbook", fake_parse)
    result = asyncio.run(
        DocumentProcessor().process_document(
            file_url="https://example.invalid/file",
            filename="annexure-4-bcci-selection-committee-presentation",
            workspace_id="ws",
            user_id="u",
            document_id="d",
            extension=extension,
        )
    )
    return seen, result


def test_process_document_routes_on_extension_but_keeps_filename(monkeypatch):
    seen, result = _run_until_routed(monkeypatch, _deck_bytes(), "pptx")
    assert seen == {"text_path": "annexure-4-bcci-selection-committee-presentation.pptx"}
    assert result["filename"] == "annexure-4-bcci-selection-committee-presentation"


def test_process_document_sends_workbooks_down_the_table_path(monkeypatch):
    # By extension, and by content when the backend sends no extension.
    for extension in ("xlsx", ""):
        seen, _ = _run_until_routed(monkeypatch, _workbook_bytes(), extension)
        assert seen == {"table_path": True}, extension
